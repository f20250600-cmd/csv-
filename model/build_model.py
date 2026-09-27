from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter as CL

F="Arial"
BLUE=Font(name=F,color="0000FF"); BLK=Font(name=F); GRN=Font(name=F,color="008000")
BOLD=Font(name=F,bold=True); TITLE=Font(name=F,bold=True,size=14); HDR=Font(name=F,bold=True,color="FFFFFF")
YEL=PatternFill("solid",fgColor="FFFF00"); HFILL=PatternFill("solid",fgColor="1F3864"); SUB=PatternFill("solid",fgColor="D9E1F2")
thin=Side(style="thin",color="BFBFBF"); BOX=Border(top=thin,bottom=thin,left=thin,right=thin)
PCT='0.0%;(0.0%);-'; USD='$#,##0;($#,##0);-'; USD2='$#,##0.00;($#,##0.00);-'; NUM='#,##0;(#,##0);-'; MULT='0.0"x"'

wb=Workbook()
def cell(ws,ref,v,font=BLK,fmt=None,fill=None,bold=False):
    c=ws[ref]; c.value=v; c.font=Font(name=F,bold=bold,color=font.color) if bold else font
    if fmt: c.number_format=fmt
    if fill: c.fill=fill
    return c
def header(ws,row,cols,texts):
    for col,t in zip(cols,texts):
        c=ws.cell(row=row,column=col,value=t); c.font=HDR; c.fill=HFILL; c.alignment=Alignment(horizontal="center",wrap_text=True,vertical="center")

TK=["C","D","E","F","G","H"]  # 6 ticker slots on Inputs / DCF

# ---------------- README ----------------
rd=wb.active; rd.title="README"
lines=[("Software Valuation Model — owner-FCF DCF with durability-weighted buy prices",TITLE),("",None),
("HOW TO UPDATE AFTER A NEW FILING",BOLD),
("1. Go to 'Inputs'. Change only BLUE cells (yellow = key judgment calls).",None),
("2. From the 10-Q / 10-K / earnings release, update: price, diluted shares, total debt, cash & investments, FCF (OCF − capex), SBC.",None),
("3. Put any one-time cash items (tax catch-ups, legal settlements, working-capital swings) in 'One-time items to remove'.",None),
("4. Revisit the judgment cells: starting growth (anchor to guidance) and P(durable).",None),
("5. Read results on 'Summary'. Use 'Sensitivity' (pick a ticker in the yellow cell) to see how r and g move value.",None),
("6. Record what changed on 'Filing Log' so you can see how your target evolved.",None),
("",None),
("ADDING A COMPANY",BOLD),
("Fill an empty column (G or H) on 'Inputs'. Everything else picks it up automatically (6 slots total).",None),
("",None),
("COLOR LEGEND",BOLD),
("Blue text = hardcoded input you edit · Black = formula (do not edit) · Green = link from another sheet · Yellow fill = key judgment assumption",None),
("",None),
("METHOD",BOLD),
("Owner FCF = normalized FCF − stock-based comp (SBC treated as a real cost).",None),
("Going-concern value: 10-year DCF, growth fades linearly from 'starting growth' to terminal growth, Gordon terminal value, minus net debt, per diluted share.",None),
("Bear value: owner FCF shrinks by the bear decline rate every year forever: OwnerFCF×(1−d)/(r+d) − net debt.",None),
("Weighted value (conservative r) = P(durable) × going-concern + (1 − P(durable)) × bear.",None),
("Buy below = MIN(weighted value, going-concern value × (1 − margin of safety)), both at the conservative discount rate.",None),
("Market-implied growth = starting growth that makes going-concern value equal the price (solved on the 'ImpliedCalc' grid, 0.5% steps, −20% to +40%).",None),
("",None),
("CAVEATS",BOLD),
("Starting values (Sep-2026) came from web-search summaries of filings/press releases, not the filings themselves — verify against 10-K/10-Q.",None),
("Cells marked EST in the source notes are analyst estimates. P(durable) is a judgment, not data. This is analysis, not personal financial advice.",None)]
for i,(t,f) in enumerate(lines,1):
    c=rd.cell(row=i,column=1,value=t); c.font=f or BLK
rd.column_dimensions["A"].width=150

# ---------------- Inputs ----------------
ip=wb.create_sheet("Inputs")
ip.column_dimensions["A"].width=44; ip.column_dimensions["B"].width=12
for c in TK: ip.column_dimensions[c].width=16
ip.column_dimensions["I"].width=60
cell(ip,"A1","Inputs — edit BLUE cells only",TITLE)
cell(ip,"A3","GLOBAL ASSUMPTIONS",BOLD)
glob=[(4,"Base discount rate",0.095,"Typical large-cap software cost of equity; CAPM with raw betas 1.3–1.5 gives ~10–11%."),
      (5,"Conservative discount rate (used for buy price)",0.105,"Higher hurdle used for weighted value and buy-below price."),
      (6,"Terminal growth",0.03,"Perpetual growth after year 10. 2% cuts value ~11% vs 3%."),
      (7,"Bear case: annual owner-FCF decline",0.05,"Structural-decline scenario (e.g., AI disruption)."),
      (8,"Margin of safety on going-concern value",0.30,"Buy-below cross-check: going-concern × (1 − this).")]
for r,l,v,n in glob:
    cell(ip,f"A{r}",l); cell(ip,f"B{r}",v,BLUE,PCT,YEL); cell(ip,f"I{r}",n,Font(name=F,italic=True,color="595959"))
cell(ip,"A10","COMPANY INPUTS",BOLD)
header(ip,11,[1,2]+list(range(3,9))+[9],["Item","Units","Slot 1","Slot 2","Slot 3","Slot 4","Slot 5","Slot 6","Notes"])
rows={}
spec=[("tick","Ticker","text"),("date","Price date","text"),("price","Share price","$"),("sh","Diluted shares","mm"),
("debt","Total debt","$mm"),("cash","Cash & investments","$mm"),("oth","Other claims (pending deals, etc.)","$mm"),
("nd","Net debt","$mm"),("fcf","Reported FCF (OCF − capex), TTM or guided","$mm"),("one","One-time items to remove","$mm"),
("int","After-tax net interest add-back (unlevered)","$mm"),("nfcf","Normalized FCF","$mm"),("sbc","Stock-based compensation","$mm"),
("own","Owner FCF (normalized FCF − SBC)","$mm"),("g","Starting growth (year 1)","%"),("p","P(durable) — judgment","%"),
("per","FCF period / basis","text"),("upd","Last updated","text"),("src","Source notes","text")]
r0=12
for i,(k,l,u) in enumerate(spec):
    r=r0+i; rows[k]=r; cell(ip,f"A{r}",l); cell(ip,f"B{r}",u)
data={
"tick":["INTU","ADBE","ADSK","CRM"],
"date":["2026-09-26"]*4,
"price":[275.79,235.47,209.40,234.02],
"sh":[267,396,214,821],
"debt":[7700,6360,3980,39500],
"cash":[7200,5640,1830,11400],
"oth":[0,0,0,1500],
"fcf":[8620,10500,2740,15050],
"one":[1120,0,0,0],
"int":[0,0,0,950],
"sbc":[1900,2100,850,3600],
"g":[0.09,0.10,0.12,0.09],
"p":[0.5,0.5,0.8,0.8],
"per":["FY26 (Jul-26) reported","TTM to Aug-26 (est.)","FY27 (Jan-27) guidance midpoint","FY27 (Jan-27) guidance (+4-5%)"],
"upd":["2026-09-27"]*4,
"src":["FY26 10-K/Q4 release: OCF $8.84B, capex $221M, cash+inv $7.2B, debt $7.7B. One-time = EST R&D-expensing tax catch-up. SBC ~$1.7-2.1B. FY27 rev guide +9-10%.",
       "Q3 FY26 10-Q (Aug-28-26): cash $4.36B + ST inv $1.28B; debt $6.36B. TTM FCF ~$10.3-10.5B. SBC ~$2.1B/yr. FY26 rev ~+12%, ARR +11.2%.",
       "Q2 FY27 10-Q: cash $4.36B, debt $2.98B; then MaintainX $3.53B paid Aug-3 with $1.0B term loan (EST post-deal). FY27 FCF guide $2.725-2.75B; SBC FY26 $788M.",
       "Q2 FY27: cash $8.31B + securities $3.09B; debt ~$39.5B after $25B ASR; Contentful $1.5B pending. FY27 FCF guide +4-5% on $14.4B. Interest add-back EST. Diluted shares 821M."]}
for k,vals in data.items():
    r=rows[k]
    for j,v in enumerate(vals):
        fmt={"price":USD2,"sh":NUM,"debt":NUM,"cash":NUM,"oth":NUM,"fcf":NUM,"one":NUM,"int":NUM,"sbc":NUM,"g":PCT,"p":PCT}.get(k)
        fill=YEL if k in ("g","p","one") else None
        cell(ip,f"{TK[j]}{r}",v,BLUE,fmt,fill)
# empty slots styled as inputs
for c in TK[4:]:
    for k,_,_ in spec:
        if k in ("nd","nfcf","own"): continue
        cc=ip[f"{c}{rows[k]}"]; cc.font=BLUE
        if k in ("g","p","one"): cc.fill=YEL
        cc.number_format={"price":USD2,"g":PCT,"p":PCT}.get(k,NUM)
for c in TK:
    t=f"{c}{rows['tick']}"
    cell(ip,f"{c}{rows['nd']}",f'=IF({t}="","",{c}{rows["debt"]}-{c}{rows["cash"]}+{c}{rows["oth"]})',BLK,NUM)
    cell(ip,f"{c}{rows['nfcf']}",f'=IF({t}="","",{c}{rows["fcf"]}-{c}{rows["one"]}+{c}{rows["int"]})',BLK,NUM)
    cell(ip,f"{c}{rows['own']}",f'=IF({t}="","",{c}{rows["nfcf"]}-{c}{rows["sbc"]})',BOLD,NUM)
    ip[f"{c}{rows['src']}"].alignment=Alignment(wrap_text=True,vertical="top")
ip.row_dimensions[rows["src"]].height=120
notes={"one":"Remove non-recurring cash (e.g., INTU FY26 R&D tax catch-up, EST $1.12B).",
"int":"Use if FCF is after big interest costs and you subtract net debt (avoids double count).",
"oth":"Pending cash acquisitions, underfunded pensions, etc.",
"g":"Anchor to revenue/FCF guidance; fades linearly to terminal growth by year 10.",
"p":"Your probability the business stays a durable going concern (vs structural decline). Biggest driver for INTU/ADBE.",
"sh":"Use latest diluted count (CRM fell 962M→821M after the ASR)."}
for k,n in notes.items(): cell(ip,f"I{rows[k]}",n,Font(name=F,italic=True,color="595959"))
ip.freeze_panes="C12"
R=rows

# ---------------- DCF ----------------
d=wb.create_sheet("DCF")
d.column_dimensions["A"].width=40; d.column_dimensions["B"].width=8
for c in TK: d.column_dimensions[c].width=14
cell(d,"A1","DCF engine — all formulas (no inputs here)",TITLE)
header(d,3,[1,2]+list(range(3,9)),["Line","Year"]+[f"Slot {i}" for i in range(1,7)])
def link(c,k): return f"Inputs!{c}{R[k]}"
cell(d,"A4","Ticker",BOLD)
for c in TK: cell(d,f"{c}4",f'=IF({link(c,"tick")}="","",{link(c,"tick")})',GRN)
cell(d,"A5","Owner FCF, year 0 ($mm)")
cell(d,"A6","Starting growth")
for c in TK:
    cell(d,f"{c}5",f'=IF({c}$4="","",{link(c,"own")})',GRN,NUM); cell(d,f"{c}6",f'=IF({c}$4="","",{link(c,"g")})',GRN,PCT)
def block(start,title):
    cell(d,f"A{start-1}",title,BOLD); 
    for c in ["A","B"]+TK: d[f"{c}{start-1}"].fill=SUB
    for t in range(1,11): cell(d,f"B{start+t-1}",t)
GR=9; block(GR,"Growth rate by year")
for t in range(10):
    r=GR+t; cell(d,f"A{r}",f"Year {t+1}")
    for c in TK: cell(d,f"{c}{r}",f'=IF({c}$4="","",{c}$6+(Inputs!$B$6-{c}$6)*($B{r}-1)/9)',BLK,PCT)
FR=21; block(FR,"Owner FCF ($mm)")
for t in range(10):
    r=FR+t; cell(d,f"A{r}",f"Year {t+1}")
    prev="$5" if t==0 else f"{r-1}"
    for c in TK:
        p=f"{c}$5" if t==0 else f"{c}{r-1}"
        cell(d,f"{c}{r}",f'=IF({c}$4="","",{p}*(1+{c}{GR+t}))',BLK,NUM)
def pvblock(start,rate,label):
    block(start,label)
    for t in range(10):
        r=start+t; cell(d,f"A{r}",f"Year {t+1}")
        for c in TK: cell(d,f"{c}{r}",f'=IF({c}$4="","",{c}{FR+t}/(1+{rate})^$B{r})',BLK,NUM)
PB=33; pvblock(PB,"Inputs!$B$4","PV of owner FCF @ base discount rate ($mm)")
PC=45; pvblock(PC,"Inputs!$B$5","PV of owner FCF @ conservative discount rate ($mm)")
S=57
cell(d,f"A{S-1}","Valuation",BOLD)
for c in ["A","B"]+TK: d[f"{c}{S-1}"].fill=SUB
lab=[("Sum PV years 1-10 (base)",NUM),("PV terminal value (base)",NUM),("Enterprise value (base)",NUM),("Going-concern value / share (base)",USD2),("Terminal value % of EV (base)",PCT),
     ("Sum PV years 1-10 (conservative)",NUM),("PV terminal value (conservative)",NUM),("Enterprise value (conservative)",NUM),("Going-concern value / share (conservative)",USD2),
     ("Bear value / share (conservative)",USD2),("Bear value / share (base)",USD2)]
for i,(l,f) in enumerate(lab): cell(d,f"A{S+i}",l,BOLD if "share" in l else BLK)
for c in TK:
    g=f'{c}$4=""'
    sh=link(c,"sh"); nd=link(c,"nd")
    fs=[f'=SUM({c}{PB}:{c}{PB+9})',
        f'={c}{FR+9}*(1+Inputs!$B$6)/(Inputs!$B$4-Inputs!$B$6)/(1+Inputs!$B$4)^10',
        f'={c}{S}+{c}{S+1}',
        f'=({c}{S+2}-{nd})/{sh}',
        f'={c}{S+1}/{c}{S+2}',
        f'=SUM({c}{PC}:{c}{PC+9})',
        f'={c}{FR+9}*(1+Inputs!$B$6)/(Inputs!$B$5-Inputs!$B$6)/(1+Inputs!$B$5)^10',
        f'={c}{S+5}+{c}{S+6}',
        f'=({c}{S+7}-{nd})/{sh}',
        f'=({c}$5*(1-Inputs!$B$7)/(Inputs!$B$5+Inputs!$B$7)-{nd})/{sh}',
        f'=({c}$5*(1-Inputs!$B$7)/(Inputs!$B$4+Inputs!$B$7)-{nd})/{sh}']
    for i,fm in enumerate(fs):
        cell(d,f"{c}{S+i}",f'=IF({g},"",{fm[1:]})',BOLD if "share" in lab[i][0] else BLK,lab[i][1])
d.freeze_panes="C4"
ROW=dict(gcb=S+3,gcc=S+8,bearc=S+9,bearb=S+10,tv=S+4)

# ---------------- ImpliedCalc ----------------
ic=wb.create_sheet("ImpliedCalc")
cell(ic,"A1","Helper: value multiple of year-0 owner FCF for each starting growth (do not edit)",TITLE)
cell(ic,"A3","Start g",BOLD)
for t in range(1,11): cell(ic,f"{CL(t+1)}3",t,BOLD)
cell(ic,"L3","Mult @ base r",BOLD); cell(ic,"M3","Mult @ cons r",BOLD)
cell(ic,"L2","Years 1-10 = cumulative growth factors",Font(name=F,italic=True))
N0=4; NR=121
for i in range(NR):
    r=N0+i
    cell(ic,f"A{r}",f"=ROUND(-0.2+{i}*0.005,4)",BLK,PCT)
    for t in range(1,11):
        col=CL(t+1); prev="1" if t==1 else f"{CL(t)}{r}"
        cell(ic,f"{col}{r}",f"={prev}*(1+$A{r}+(Inputs!$B$6-$A{r})*({col}$3-1)/9)",BLK,"0.000")
    for col,rate in (("L","Inputs!$B$4"),("M","Inputs!$B$5")):
        cell(ic,f"{col}{r}",f"=SUMPRODUCT(B{r}:K{r},1/(1+{rate})^$B$3:$K$3)+K{r}*(1+Inputs!$B$6)/({rate}-Inputs!$B$6)/(1+{rate})^10",BLK,"0.00")
LAST=N0+NR-1
ic.column_dimensions["A"].width=10

# ---------------- Summary ----------------
sm=wb.create_sheet("Summary",1)
cell(sm,"A1","Summary — price targets (all formulas; edit 'Inputs')",TITLE)
cell(sm,"A2",'="Base r "&TEXT(Inputs!B4,"0.0%")&" · Conservative r "&TEXT(Inputs!B5,"0.0%")&" · Terminal g "&TEXT(Inputs!B6,"0.0%")&" · Bear decline "&TEXT(Inputs!B7,"0.0%")&" · MOS "&TEXT(Inputs!B8,"0%")',Font(name=F,italic=True))
cols=["Slot","Ticker","Price","EV / owner FCF","SBC % of FCF","Starting growth","Going-concern value (base r)","Going-concern value (cons. r)",
"Bear value (cons. r)","P(durable)","Weighted value (cons. r)","Going-concern × (1−MOS)","BUY BELOW","Price vs buy-below","Implied growth (base r)","Implied growth (cons. r)","Verdict"]
header(sm,4,range(1,len(cols)+1),cols); sm.row_dimensions[4].height=45
for i,w in enumerate([6,9,11,11,11,11,14,14,13,11,14,14,13,12,13,13,12],1): sm.column_dimensions[CL(i)].width=w
IDX=lambda sheet,row,slot: f"INDEX({sheet}!$C${row}:$H${row},{slot})"
for s in range(1,7):
    r=4+s; A=f"$A{r}"
    cell(sm,f"A{r}",s)
    tk=IDX("Inputs",R["tick"],A)
    cell(sm,f"B{r}",f'=IF({tk}="","",{tk})',BOLD)
    g=f'$B{r}=""'
    def put(col,expr,fmt,font=BLK):
        cell(sm,f"{col}{r}",f'=IF({g},"",{expr})',font,fmt)
    price=IDX("Inputs",R["price"],A); sh=IDX("Inputs",R["sh"],A); nd=IDX("Inputs",R["nd"],A); own=IDX("Inputs",R["own"],A)
    put("C",price,USD2,GRN)
    put("D",f"({price}*{sh}+{nd})/{own}",MULT)
    put("E",f'{IDX("Inputs",R["sbc"],A)}/{IDX("Inputs",R["nfcf"],A)}',PCT)
    put("F",IDX("Inputs",R["g"],A),PCT,GRN)
    put("G",IDX("DCF",ROW["gcb"],A),USD)
    put("H",IDX("DCF",ROW["gcc"],A),USD)
    put("I",IDX("DCF",ROW["bearc"],A),USD)
    put("J",IDX("Inputs",R["p"],A),PCT,GRN)
    put("K",f"J{r}*H{r}+(1-J{r})*I{r}",USD)
    put("L",f"H{r}*(1-Inputs!$B$8)",USD)
    put("M",f"MIN(K{r},L{r})",USD,BOLD)
    put("N",f"M{r}/C{r}-1",PCT)
    for col,mc in (("O","L"),("P","M")):
        need=f"({price}*{sh}+{nd})/{own}"
        put(col,f'IF({need}<ImpliedCalc!${mc}${N0},"< -20%",IF({need}>ImpliedCalc!${mc}${LAST},"> 40%",INDEX(ImpliedCalc!$A${N0}:$A${LAST},MATCH({need},ImpliedCalc!${mc}${N0}:${mc}${LAST},1))))',PCT)
    put("Q",f'IF(C{r}<=M{r},"BUY ZONE",IF(C{r}<=M{r}*1.1,"WATCH","NO BUY"))',None,BOLD)
    for c in range(1,18): sm.cell(row=r,column=c).border=BOX
sm["M4"].fill=PatternFill("solid",fgColor="C00000")
sm.conditional_formatting.add("Q5:Q10",CellIsRule(operator="equal",formula=['"BUY ZONE"'],fill=PatternFill("solid",fgColor="C6EFCE")))
sm.conditional_formatting.add("Q5:Q10",CellIsRule(operator="equal",formula=['"WATCH"'],fill=PatternFill("solid",fgColor="FFEB9C")))
sm.conditional_formatting.add("Q5:Q10",CellIsRule(operator="equal",formula=['"NO BUY"'],fill=PatternFill("solid",fgColor="FFC7CE")))
cell(sm,"A12","Buy below = MIN(weighted value, going-concern × (1 − MOS)), both at the conservative discount rate. Verdict: BUY ZONE if price ≤ buy-below; WATCH if within 10% above; else NO BUY.",Font(name=F,italic=True))
cell(sm,"A13","Implied growth = starting growth at which going-concern value equals today's price (0.5% steps). Compare with guidance: far above guidance = market optimistic; below = market pessimistic.",Font(name=F,italic=True))
cell(sm,"A15","P(durable) sensitivity of weighted value (cons. r)",BOLD)
ps=[0.3,0.5,0.7,0.9]
header(sm,16,range(2,7),["Ticker"]+[None]*4)
for j,p in enumerate(ps):
    c=sm.cell(row=16,column=3+j,value=p); c.number_format='"P = "0%'; c.font=Font(name=F,color="0000FF",bold=True); c.fill=YEL
for s in range(1,7):
    r=16+s
    cell(sm,f"B{r}",f"=B{4+s}",GRN)
    for j in range(4):
        col=CL(3+j)
        cell(sm,f"{col}{r}",f'=IF($B{r}="","",{col}$16*$H{4+s}+(1-{col}$16)*$I{4+s})',BLK,USD)
sm.freeze_panes="C5"

# ---------------- Sensitivity ----------------
se=wb.create_sheet("Sensitivity",2)
cell(se,"A1","Sensitivity — going-concern value per share",TITLE)
cell(se,"A3","Pick ticker:",BOLD); cell(se,"B3","INTU",BLUE,None,YEL)
dv=DataValidation(type="list",formula1="=Inputs!$C$12:$H$12",allow_blank=False); se.add_data_validation(dv); dv.add("B3")
cell(se,"A4","Slot"); cell(se,"B4",'=MATCH(B3,Inputs!$C$12:$H$12,0)')
cell(se,"A5","Price"); cell(se,"B5",f'=INDEX(Inputs!$C${R["price"]}:$H${R["price"]},B4)',GRN,USD2)
cell(se,"A6","Owner FCF ($mm)"); cell(se,"B6",f'=INDEX(Inputs!$C${R["own"]}:$H${R["own"]},B4)',GRN,NUM)
cell(se,"A7","Net debt ($mm)"); cell(se,"B7",f'=INDEX(Inputs!$C${R["nd"]}:$H${R["nd"]},B4)',GRN,NUM)
cell(se,"A8","Shares (mm)"); cell(se,"B8",f'=INDEX(Inputs!$C${R["sh"]}:$H${R["sh"]},B4)',GRN,NUM)
cell(se,"A10","Discount rate ↓ / Starting growth →  (edit the blue/yellow cells)",BOLD)
gs=[-0.05,0,0.05,0.09,0.13,0.17]; rs=[0.085,0.095,0.105,0.115,0.125]
for j,g in enumerate(gs): cell(se,f"{CL(2+j)}11",g,BLUE,PCT,YEL)
cell(se,"A25","Helper: cumulative growth factors per column (do not edit)",BOLD)
cell(se,"A26","Year",BOLD)
for j in range(len(gs)): cell(se,f"{CL(2+j)}26",f"={CL(2+j)}11",GRN,PCT)
for t in range(1,11):
    r=26+t; cell(se,f"A{r}",t)
    for j in range(len(gs)):
        col=CL(2+j); prev="1" if t==1 else f"{col}{r-1}"
        cell(se,f"{col}{r}",f"={prev}*(1+{col}$11+(Inputs!$B$6-{col}$11)*($A{r}-1)/9)",BLK,"0.000")
for i,rv in enumerate(rs):
    r=12+i; cell(se,f"A{r}",rv,BLUE,PCT,YEL)
    for j in range(len(gs)):
        col=CL(2+j)
        mult=f"SUMPRODUCT({col}$27:{col}$36,1/(1+$A{r})^$A$27:$A$36)+{col}$36*(1+Inputs!$B$6)/($A{r}-Inputs!$B$6)/(1+$A{r})^10"
        cell(se,f"{col}{r}",f'=IF($B$4="","",($B$6*({mult})-$B$7)/$B$8)',BLK,USD)
        se[f"{col}{r}"].border=BOX
se.conditional_formatting.add("B12:G16",FormulaRule(formula=["B12>=$B$5"],fill=PatternFill("solid",fgColor="C6EFCE")))
se.conditional_formatting.add("B12:G16",FormulaRule(formula=["B12<$B$5"],fill=PatternFill("solid",fgColor="FFC7CE")))
cell(se,"A18","Green = value above today's price; red = below. Terminal growth comes from Inputs!B6; change it there to test 2% vs 3%.",Font(name=F,italic=True))
se.column_dimensions["A"].width=22
for j in range(6): se.column_dimensions[CL(2+j)].width=12

# ---------------- Filing Log ----------------
lg=wb.create_sheet("Filing Log")
cell(lg,"A1","Filing log — record every update",TITLE)
hd=["Date","Ticker","Filing / event","What changed (inputs)","Old buy-below","New buy-below","Price on date","Thesis status / notes"]
header(lg,3,range(1,9),hd)
ex=["2026-09-27","ALL","Initial build (Q2/Q3 filings, Sep-26 prices)","Baseline inputs; INTU FCF normalized for tax catch-up; CRM post-ASR shares/debt; ADSK post-MaintainX debt",None,None,None,"INTU watch ~$245; ADBE buy small ~$255; ADSK/CRM no buy"]
for j,v in enumerate(ex):
    c=lg.cell(row=4,column=j+1,value=v); c.font=BLUE
for j,w in enumerate([12,8,34,60,13,13,12,50],1): lg.column_dimensions[CL(j)].width=w
cell(lg,"A6","Tip: copy the Summary 'BUY BELOW' value as a number into 'New buy-below' each time, so history is preserved.",Font(name=F,italic=True))

wb.calculation.fullCalcOnLoad=True
wb.save("/home/user/csv-/model/Valuation_Model.xlsx")

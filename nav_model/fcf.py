"""
"NAV + FCF generated" layer, built so nothing is counted twice.

Today's NAV is already the PV of every future year of asset cash flow, so today's NAV + cumulative FCF
double counts the FCF years. The consistent version is a forward NAV:

    Value per share at end-2030 = cumulative equity FCF 2027-2030 (after overhead, interest, cash tax,
                                  preferred dividends; includes retail/platform cash)
                                + equity NAV at end-2030 (assets valued only on cash flows from 2031 onward).

Compared with today's price, this gives an implied annual return. Discounted back at the equity rate, it should
land near today's NAV incl. platform, which is the reconciliation check.
"""
import csv
import os

import inputs as I
import nav_model as N

YEARS = list(range(I.FIRST_YEAR, I.FCF_HORIZON_END + 1))


def annual_equity_fcf(co, R):
    """Per-year equity FCF build ($bn, nominal) for one company/scenario result R."""
    fin = I.FINANCING[co]
    ren_value = sum(r["pv_mm"] for r in R["rows"] if r["tech"] == "renewable") / 1000
    out = []
    for y in YEARS:
        infl = N.Deck.infl(y)
        streams = {k: N.annual_stream(R, y, k) for k in N.STREAM_KEYS}
        revenue = sum(streams[k] for k in ("energy", "capacity", "ppa", "zec"))
        asset_margin = sum(streams.values()) + ren_value * I.RENEWABLE_CASH_YIELD
        ptc = streams["ptc45u"]
        platform = I.PLATFORM[co]["ebitda_mm"] / 1000 * infl
        overhead = I.CORP_OVERHEAD_MM[co] / 1000 * infl
        interest = fin["net_debt_bn"] * fin["cost_of_debt"]
        pref = fin["pref_bn"] * fin["pref_rate"]
        taxable = asset_margin - ptc + platform - overhead - interest
        tax = I.CASH_TAX_RATE * max(taxable, 0.0)
        fcf = asset_margin + platform - overhead - interest - tax - pref
        out.append({"year": y, "revenue": revenue, "asset_margin": asset_margin, "platform": platform,
                    "overhead": overhead, "interest": interest, "tax": tax, "pref": pref, "fcf": fcf})
    return out


def forward_nav(co, scen):
    """Equity NAV at end of FCF horizon: assets valued on cash flows after the horizon only ($bn)."""
    saved = (I.VALUATION_DATE, I.FIRST_YEAR)
    try:
        I.VALUATION_DATE = I.FCF_HORIZON_END + 1.0
        I.FIRST_YEAR = I.FCF_HORIZON_END + 1
        R = N.company_nav(co, N.Deck(scen))
    finally:
        I.VALUATION_DATE, I.FIRST_YEAR = saved
    platform = R["platform"] * N.Deck.infl(I.FCF_HORIZON_END)
    return {"gav": R["gav"], "overhead": R["overhead"], "net_claims": R["net_claims"],
            "equity_gen": R["equity"], "equity": R["equity"] + platform}


def run(results, mkt, out_dir):
    md = ["### NAV + FCF generated (forward-NAV method, no double counting)\n",
          f"Equity FCF {YEARS[0]}-{YEARS[-1]} (after overhead, interest, 20% cash tax and preferred dividends; includes retail "
          f"platform cash; no growth capex, no buybacks) + equity NAV at end-{YEARS[-1]} of the remaining asset life. "
          f"Implied return = (value at end-{YEARS[-1]} / price)^(1/{YEARS[-1] + 1 - I.VALUATION_DATE:.2f}) - 1.\n"]
    rows_csv, summary = [], []
    horizon = YEARS[-1] + 1 - I.VALUATION_DATE
    for co in ("CEG", "VST"):
        m = mkt[co]
        sh = results[co]["Base"]["shares"]
        md.append(f"#### {co} (price ${m['price']:.2f}, market cap ${m['mcap']:.1f}bn)\n")
        md.append("| $ per share unless noted | Low | Base | High | Extreme |")
        md.append("|---|---|---|---|---|")
        per = {}
        for s in I.SCENARIO_ORDER:
            fcfs = annual_equity_fcf(co, results[co][s])
            cum = sum(f["fcf"] for f in fcfs)
            pv_fcf = sum(f["fcf"] * (1 + I.EQUITY_DISCOUNT_RATE) ** -(f["year"] + 0.5 - I.VALUATION_DATE) for f in fcfs)
            fwd = forward_nav(co, s)
            total = cum + fwd["equity"]
            per[s] = {
                "fcf27": fcfs[0]["fcf"], "yield27": fcfs[0]["fcf"] / m["mcap"],
                "cum_ps": cum / sh * 1000, "fwd_ps": fwd["equity"] / sh * 1000, "total_ps": total / sh * 1000,
                "irr": (total / sh * 1000 / m["price"]) ** (1 / horizon) - 1 if total > 0 else None,
                "pv_total_ps": (pv_fcf + fwd["equity"] * (1 + I.EQUITY_DISCOUNT_RATE) ** -horizon) / sh * 1000,
                "today_ps": results[co][s]["nav_ps_plat"],
                "naive_ps": results[co][s]["nav_ps_plat"] + cum / sh * 1000,
            }
            for f in fcfs:
                rows_csv.append({"company": co, "scenario": s, **{k: round(v, 3) if k != "year" else v for k, v in f.items()}})
            summary.append({"company": co, "scenario": s, **{k: (round(v, 4) if v is not None else None) for k, v in per[s].items()}})
        L = lambda label, key, f: md.append(f"| {label} | " + " | ".join(f(per[s][key]) for s in I.SCENARIO_ORDER) + " |")
        L(f"Equity FCF {YEARS[0]} ($bn)", "fcf27", lambda v: f"{v:.1f}")
        L(f"FCF yield on market cap, {YEARS[0]}", "yield27", lambda v: f"{v:.1%}")
        L(f"(a) Cumulative equity FCF {YEARS[0]}-{YEARS[-1]}", "cum_ps", lambda v: f"{v:,.0f}")
        L(f"(b) Equity NAV at end-{YEARS[-1]} (incl. platform)", "fwd_ps", lambda v: f"{v:,.0f}")
        L(f"**(a)+(b) Value per share at end-{YEARS[-1]}**", "total_ps", lambda v: f"**{v:,.0f}**")
        L(f"**Implied annual return from ${m['price']:.2f}**", "irr", lambda v: "n.m." if v is None else f"**{v:+.1%}**")
        L("Check: (a)+(b) discounted to today at 9%", "pv_total_ps", lambda v: f"{v:,.0f}")
        L("Today's NAV incl. platform (from main model)", "today_ps", lambda v: f"{v:,.0f}")
        L("Naive today's NAV + cumulative FCF (**double counts - do not use**)", "naive_ps", lambda v: f"~~{v:,.0f}~~")
        md.append("")
        md.append(f"| {co} annual equity FCF build, Base ($bn) | " + " | ".join(str(y) for y in YEARS) + " |")
        md.append("|---|" + "---|" * len(YEARS))
        fb = annual_equity_fcf(co, results[co]["Base"])
        for key, label, sign in (("revenue", "Revenue (energy + capacity + PPA + ZEC)", 1), ("asset_margin", "Asset cash margin", 1),
                                 ("platform", "+ Retail/platform EBITDA", 1), ("overhead", "- Corporate overhead", -1),
                                 ("interest", "- Interest", -1), ("tax", "- Cash tax", -1), ("pref", "- Preferred dividends", -1),
                                 ("fcf", "**= Equity FCF**", 1)):
            md.append(f"| {label} | " + " | ".join(f"{sign * f[key]:.2f}" for f in fb) + " |")
        md.append("")
    with open(os.path.join(out_dir, "fcf_nav.md"), "w") as f:
        f.write("\n".join(md))
    with open(os.path.join(out_dir, "equity_fcf_annual.csv"), "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(rows_csv[0]))
        w.writeheader()
        w.writerows(rows_csv)
    with open(os.path.join(out_dir, "nav_plus_fcf_summary.csv"), "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(summary[0]))
        w.writeheader()
        w.writerows(summary)
    try:
        from openpyxl import load_workbook
        path = os.path.join(out_dir, "nav_model.xlsx")
        wb = load_workbook(path)
        for name in ("NAV + FCF", "Equity FCF annual"):
            if name in wb.sheetnames:
                del wb[name]
        ws = wb.create_sheet("NAV + FCF")
        ws.append(list(summary[0]))
        for r in summary:
            ws.append(list(r.values()))
        ws = wb.create_sheet("Equity FCF annual")
        ws.append(list(rows_csv[0]))
        for r in rows_csv:
            ws.append(list(r.values()))
        wb.save(path)
    except ImportError:
        pass
    return summary

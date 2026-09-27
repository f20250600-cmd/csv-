"""
Replacement-cost theory + FCF.

    Asset value     = depreciated replacement cost (DRC): what it would cost to build the same fleet today
                      (2026 $/kW new-build) x MW x remaining-life fraction. Uses no earnings or power prices.
    Equity value    = DRC - net debt & other claims
                    + PV of equity FCF the existing fleet generates over the replacement lead time
                      (a rival cannot bring equivalent capacity online sooner, so that cash is a rent on top of cost).

Two bases for nuclear: like-for-like (new AP1000) and functional (new CCGT per accredited MW).
"""
import csv
import os

import inputs as I
import nav_model as N
import fcf as F


def _lookup(table, name):
    for k, v in table.items():
        if name.startswith(k) or k in name:
            return v
    return None


def asset_drc(a, basis):
    tech = a["tech"]
    if tech == "nuclear":
        rc = I.REPLACEMENT_COST_KW["nuclear_like"] if basis == "like" else I.NUCLEAR_FUNCTIONAL_KW
    else:
        rc = I.REPLACEMENT_COST_KW[tech]
    frac = _lookup(I.LIFE_FRAC_OVERRIDE, a["name"])
    if frac is None:
        cod = _lookup(I.ASSET_COD, a["name"])
        end = a["end"]
        frac = max(0.0, min(1.0, (end + 1 - I.VALUATION_DATE) / (end + 1 - cod)))
    return {"name": a["name"], "tech": tech, "mw": a["mw"], "rc_kw": rc, "frac": frac,
            "rcn_bn": a["mw"] * rc / 1e6, "drc_bn": a["mw"] * rc * frac / 1e6}


def fcf_pv(co, R, horizon):
    if horizon == 0:
        return 0.0, 0.0
    years = list(range(I.FIRST_YEAR, I.FIRST_YEAR + horizon))
    fl = F.annual_equity_fcf(co, R, years)
    cum = sum(f["fcf"] for f in fl)
    pv = sum(f["fcf"] * (1 + I.EQUITY_DISCOUNT_RATE) ** -(f["year"] + 0.5 - I.VALUATION_DATE) for f in fl)
    return cum, pv


def run(results, mkt, out_dir):
    md = ["# Replacement cost + FCF\n",
          "Value/share = [depreciated replacement cost of the fleet - net debt & claims + PV (9%) of equity FCF generated over "
          "the replacement lead time] / shares. Replacement cost uses 2026 new-build $/kW and remaining-life fractions only - "
          "no earnings or power prices. FCF (after overhead, interest, tax, preferred; includes retail cash) is the only "
          "scenario-dependent part.\n"]
    rows_csv, summary = [], []
    for co in ("CEG", "VST"):
        assets = I.CEG_ASSETS if co == "CEG" else I.VST_ASSETS
        m = mkt[co]
        B = results[co]["Base"]
        sh, claims = B["shares"], B["net_claims"]
        drc = {b: [asset_drc(a, b) for a in assets] for b in ("like", "func")}
        md.append(f"## {co} (price ${m['price']:.2f}; net claims ${claims:.1f}bn; {sh:.1f}mm shares)\n")
        md.append("### Depreciated replacement cost by asset group\n")
        md.append("| Asset group | MW | New-build $/kW | Remaining-life fraction | Replacement new $bn | **DRC $bn** |")
        md.append("|---|---|---|---|---|---|")
        for r in drc["like"]:
            md.append(f"| {r['name']} | {r['mw']:,} | {r['rc_kw']:,.0f} | {r['frac']:.0%} | {r['rcn_bn']:.1f} | **{r['drc_bn']:.1f}** |")
            rows_csv.append({"company": co, "basis": "like-for-like nuclear", **{k: (round(v, 3) if isinstance(v, float) else v) for k, v in r.items()}})
        for r in drc["func"]:
            rows_csv.append({"company": co, "basis": "functional nuclear (CCGT)", **{k: (round(v, 3) if isinstance(v, float) else v) for k, v in r.items()}})
        tot = {b: (sum(r["rcn_bn"] for r in drc[b]), sum(r["drc_bn"] for r in drc[b])) for b in drc}
        nuc = {b: sum(r["drc_bn"] for r in drc[b] if r["tech"] == "nuclear") for b in drc}
        md.append(f"| **Total - nuclear rebuilt as new nuclear** | {sum(r['mw'] for r in drc['like']):,} | | | {tot['like'][0]:.1f} | **{tot['like'][1]:.1f}** |")
        md.append(f"| **Total - nuclear replaced by new gas (functional, ${I.NUCLEAR_FUNCTIONAL_KW:,.0f}/kW)** | | | | {tot['func'][0]:.1f} | **{tot['func'][1]:.1f}** |")
        md.append("")

        for b, label in (("like", "nuclear rebuilt as new nuclear ($12,100/kW)"), ("func", f"nuclear replaced by new gas (${I.NUCLEAR_FUNCTIONAL_KW:,.0f}/kW)")):
            md.append(f"### Value per share - {label}\n")
            eq_drc = tot[b][1] - claims
            md.append(f"DRC ${tot[b][1]:.1f}bn - net claims ${claims:.1f}bn = equity replacement value **${eq_drc:.1f}bn = ${eq_drc/sh*1000:,.0f}/share** (before FCF).\n")
            md.append("| FCF years added | " + " | ".join(I.SCENARIO_ORDER) + " |")
            md.append("|---|" + "---|" * len(I.SCENARIO_ORDER))
            for H in I.RC_FCF_HORIZONS:
                vals = []
                for s in I.SCENARIO_ORDER:
                    cum, pv = fcf_pv(co, results[co][s], H)
                    v = (eq_drc + pv) / sh * 1000
                    vals.append(v)
                    summary.append({"company": co, "basis": b, "horizon_yrs": H, "scenario": s, "drc_bn": round(tot[b][1], 2),
                                    "pv_fcf_bn": round(pv, 2), "cum_fcf_bn": round(cum, 2), "value_ps": round(v, 1),
                                    "price": m["price"], "price_vs_value": round(m["price"] / v - 1, 3) if v > 0 else None})
                bold = H == I.RC_HEADLINE_HORIZON
                md.append(f"| {'**' if bold else ''}{H}{' (headline)**' if bold else ''} | "
                          + " | ".join(f"{'**' if bold else ''}${v:,.0f}{'**' if bold else ''}" for v in vals) + " |")
            md.append(f"\nCurrent price ${m['price']:.2f}.\n")

        # market-implied nuclear replacement cost with headline FCF (Base)
        _, pv5 = fcf_pv(co, results[co]["Base"], I.RC_HEADLINE_HORIZON)
        non_nuc = tot["like"][1] - nuc["like"]
        nuc_mw_frac = sum(r["mw"] * r["frac"] for r in drc["like"] if r["tech"] == "nuclear")
        need = m["price"] * sh / 1000 - pv5 + claims - non_nuc
        implied_kw = need * 1e6 / nuc_mw_frac
        md.append("### Cross-checks\n")
        md.append("| | Value |")
        md.append("|---|---|")
        md.append(f"| Market EV / DRC (nuclear as new nuclear) | {m['ev']/tot['like'][1]:.2f}x |")
        md.append(f"| Market EV / DRC (nuclear as new gas) | {m['ev']/tot['func'][1]:.2f}x |")
        for s in I.SCENARIO_ORDER:
            md.append(f"| Tobin's q = cash-flow GAV / DRC (new-nuclear basis), {s} | {results[co][s]['gav']/tot['like'][1]:.2f}x |")
        md.append(f"| Nuclear new-build $/kW the price implies (Base, {I.RC_HEADLINE_HORIZON}-yr FCF, same age haircut) | ${implied_kw:,.0f}/kW |")
        md.append("")

    with open(os.path.join(out_dir, "replacement_cost.md"), "w") as f:
        f.write("\n".join(md))
    for name, rows in (("replacement_cost_assets.csv", rows_csv), ("replacement_cost_plus_fcf.csv", summary)):
        with open(os.path.join(out_dir, name), "w", newline="") as f:
            w = csv.DictWriter(f, fieldnames=list(rows[0]))
            w.writeheader()
            w.writerows(rows)
    try:
        from openpyxl import load_workbook
        path = os.path.join(out_dir, "nav_model.xlsx")
        wb = load_workbook(path)
        for name, rows in (("Replacement cost", rows_csv), ("Replacement + FCF", summary)):
            if name in wb.sheetnames:
                del wb[name]
            ws = wb.create_sheet(name)
            ws.append(list(rows[0]))
            for r in rows:
                ws.append(list(r.values()))
        wb.save(path)
    except ImportError:
        pass
    return summary

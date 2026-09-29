"""
Build a fully formula-driven Excel version of the CEG / VST asset NAV model.

Output: outputs/CEG_VST_NAV_Model.xlsx
Every result recalculates in Excel when an input changes (blue cells / yellow levers).
Defaults are taken from inputs.py so the workbook starts from the same assumptions as the Python model.
"""
import os

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

import inputs as I

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "outputs", "CEG_VST_NAV_Model.xlsx")
YEARS = list(range(2027, 2066))
Y0 = 7                                   # column G = 2027
YC = [L(Y0 + i) for i in range(len(YEARS))]
YL, YR = YC[0], YC[-1]                   # G .. AS
HUBS = ["PJMW", "PJMA", "NIHUB", "MISO", "NY", "NE", "WEST", "ERCOT"]
PJM_HUBS = {"PJMW", "PJMA", "NIHUB"}

BLUE, BLACK, GREEN = "0000FF", "000000", "008000"
F = lambda color=BLACK, bold=False, italic=False, size=10: Font(name="Arial", size=size, color=color, bold=bold, italic=italic)
YELLOW = PatternFill("solid", fgColor="FFFF00")
HDR = PatternFill("solid", fgColor="DDE6F0")
SEC = PatternFill("solid", fgColor="F2F2F2")
THIN = Border(bottom=Side(style="thin", color="999999"))
USD = '#,##0;(#,##0);"-"'
USD1 = '#,##0.0;(#,##0.0);"-"'
USD2 = '#,##0.00;(#,##0.00);"-"'
PCT = '0.0%;(0.0%);"-"'
MULT = '0.00"x"'
YRFMT = "0"

wb = Workbook()
names = {}


def name(nm, ref):
    names[nm] = ref
    wb.defined_names[nm] = DefinedName(nm, attr_text=ref)


def put(ws, cell, value, color=BLACK, fmt=None, bold=False, fill=None, note=None, italic=False):
    c = ws[cell]
    c.value = value
    c.font = F(color, bold, italic)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if note:
        c.comment = Comment(note, "model")
    return c


def header(ws, row, labels, start_col=1):
    for i, t in enumerate(labels):
        c = ws.cell(row=row, column=start_col + i, value=t)
        c.font = F(bold=True)
        c.fill = HDR
        c.alignment = Alignment(wrap_text=True, vertical="center")


def section(ws, row, text, width=6):
    for col in range(1, width + 1):
        ws.cell(row=row, column=col).fill = SEC
    c = ws.cell(row=row, column=1, value=text)
    c.font = F(bold=True, size=11)


# =====================================================================================
# README
# =====================================================================================
rd = wb.active
rd.title = "README"
lines = [
    ("CEG / VST asset-based NAV model - editable", True),
    ("", False),
    ("How to use", True),
    ("1. Change inputs on the 'Inputs' sheet (control panel at the top) and asset-level inputs on 'CEG_Assets' / 'VST_Assets'.", False),
    ("2. Everything else is formulas: price deck ('Prices'), annual asset cash flows ('CEG_Model' / 'VST_Model') and valuation ('Summary').", False),
    ("3. Read the new price target, NAV, FCF and replacement-cost value on 'Summary'.", False),
    ("", False),
    ("Colour legend", True),
    ("Blue text = hard-coded input you can change", False),
    ("Yellow fill = key lever (scenario, price/gas shifts, revenue/cost/margin adjustments, FCF haircut, price-target weights)", False),
    ("Green text = link to another sheet (a default you can overwrite with your own number)", False),
    ("Black text = formula - do not overwrite", False),
    ("", False),
    ("Main levers (Inputs sheet)", True),
    ("Scenario (Low/Base/High/Extreme) - sets long-run heat rates, capacity prices and peaker scarcity.", False),
    ("Long-run power price shift ($/MWh, 2026$) and Henry Hub shift ($/MMBtu) - move every hub.", False),
    ("Revenue / fuel & variable cost / fixed O&M / capex adjustments (%) per company - flex revenue, gross margin and costs across all assets.", False),
    ("FCF haircut (%) for years up to the horizon - stress near-term cash flow.", False),
    ("Nuclear replacement-cost basis - gas-equivalent, gas + clean premium, NOAK AP1000 or next AP1000.", False),
    ("Price-target weights - blend of NAV, forward value (FCF + NAV at horizon) and replacement cost + FCF.", False),
    ("", False),
    ("Units", True),
    ("$mm unless stated; per-share values in $; power prices $/MWh nominal (long-run inputs in 2026$); capacity $/MW-day.", False),
    ("", False),
    ("Method notes", True),
    ("Each asset: revenue (merchant energy, capacity, long-term PPA, ZEC) - fuel & variable cost = gross margin; - fixed O&M = EBITDA; - capex = pre-tax cash flow;", False),
    ("after 20% cash tax (+ untaxed 45U credit), discounted at the asset's rate to 30-Sep-2026. Sum of assets = gross asset value; less overhead and claims = equity NAV.", False),
    ("Forward value = equity FCF to the horizon + equity NAV of remaining life at the horizon. Replacement = depreciated new-build cost - claims + PV of N years FCF.", False),
    ("Sources and caveats: see REPORT.md in the same folder. Plant-group MW splits, PPA prices and several liabilities are estimates.", False),
]
for i, (t, b) in enumerate(lines, 1):
    c = rd.cell(row=i, column=1, value=t)
    c.font = F(bold=b, size=12 if i == 1 else 10)
rd["A9"].font = F(BLUE)
rd["A10"].fill = YELLOW
rd["A11"].font = F(GREEN)
rd.column_dimensions["A"].width = 150

# =====================================================================================
# INPUTS
# =====================================================================================
ws = wb.create_sheet("Inputs")
ws.column_dimensions["A"].width = 52
for col in "BCDEFGHIJKLMNOPQRSTU":
    ws.column_dimensions[col].width = 13
ws.column_dimensions["G"].width = 60
r = 1
put(ws, "A1", "INPUTS - blue = input, yellow = key lever", bold=True)
r = 3
section(ws, r, "Control panel", 7)
r += 1


def scalar(label, nm, value, fmt=None, note=None, lever=False, src=""):
    global r
    put(ws, f"A{r}", label)
    put(ws, f"B{r}", value, BLUE, fmt, fill=YELLOW if lever else None, note=note)
    if src:
        put(ws, f"G{r}", src, italic=True)
    name(nm, f"Inputs!$B${r}")
    r += 1


scalar("Scenario (Low / Base / High / Extreme)", "Scenario", "Base", lever=True, src="Selects heat rates, capacity prices, peaker scarcity from the scenario table")
dv = DataValidation(type="list", formula1='"Low,Base,High,Extreme"', allow_blank=False)
ws.add_data_validation(dv)
dv.add(f"B{r-1}")
scalar("ERCOT (Texas) scenario - 'Same' follows the main scenario", "ERCOTScen", "Same", lever=True,
       src="Lets ERCOT tighten (or loosen) independently of PJM: sets ERCOT heat rate and ERCOT peaker scarcity")
dv3 = DataValidation(type="list", formula1='"Same,Low,Base,High,Extreme"', allow_blank=False)
ws.add_data_validation(dv3)
dv3.add(f"B{r-1}")
scalar("Long-run power price shift ($/MWh, 2026$, all hubs)", "PriceShift", 0.0, USD1, lever=True, src="Added to every hub's long-run ATC from 2030 (half in 2029)")
scalar("Henry Hub shift ($/MMBtu, all years)", "GasShift", 0.0, USD2, lever=True, src="Moves gas costs AND power prices (price = heat rate x gas)")
scalar("PJM capacity price override ($/MW-day, 2026$; blank = scenario)", "PJMCapOverride", None, USD, lever=True)
scalar("Nuclear replacement-cost basis (1-4, see table below)", "NucBasis", 2, lever=True, src="1 gas-equivalent / 2 gas + clean premium / 3 NOAK AP1000 / 4 EPRI next AP1000")
dv2 = DataValidation(type="whole", operator="between", formula1="1", formula2="4")
ws.add_data_validation(dv2)
dv2.add(f"B{r-1}")
scalar("FCF / forward-value horizon (last year of FCF)", "Horizon", I.FCF_HORIZON_END, YRFMT, lever=True)
scalar("Replacement-cost FCF years added (build lead time)", "RCYears", I.RC_HEADLINE_HORIZON, "0", lever=True, src="Years of FCF added on top of depreciated replacement cost")
scalar("Target return for entry price (fat pitch)", "TargetReturn", 0.15, PCT, lever=True)
scalar("Equity discount rate (price-target forward leg, replacement FCF)", "EquityDisc", I.EQUITY_DISCOUNT_RATE, PCT)
r += 1
section(ws, r, "Global economics", 7)
r += 1
scalar("Valuation date (fractional year)", "ValDate", I.VALUATION_DATE, "0.00", src="30-Sep-2026")
scalar("Inflation / escalator", "Infl", I.INFLATION, PCT)
scalar("Cash tax rate on asset cash flow", "TaxRate", I.CASH_TAX_RATE, PCT, src="21% federal + state less depreciation shield (assumption)")
scalar("Henry Hub 2027 (2026$/MMBtu)", "HH27", I.HENRY_HUB_REAL[2027], USD2, src="Between Sep-26 prompt $2.79 and EIA STEO ~$4.60 for 2027")
scalar("Henry Hub 2028 (2026$/MMBtu)", "HH28", I.HENRY_HUB_REAL[2028], USD2)
scalar("Henry Hub long run (2026$/MMBtu)", "HHLR", I.HENRY_HUB_LR_REAL, USD2)
scalar("Delivered-fuel adder for gas plants ($/MMBtu)", "FuelAdder", I.FUEL_ADDER, USD2, src="Transport + winter fuel/power mismatch (estimate)")
scalar("45U max credit ($/MWh, 2024$)", "PTCCredit", I.PTC45U["credit"], USD2, src="CRS IN12557")
scalar("45U gross-receipts threshold ($/MWh, 2024$)", "PTCThresh", I.PTC45U["threshold"], USD2)
scalar("45U base year", "PTCBase", I.PTC45U["base_year"], YRFMT)
scalar("45U last year", "PTCLast", I.PTC45U["last_year"], YRFMT)
scalar("Corporate overhead discount rate", "OHDisc", 0.085, PCT)
scalar("Corporate overhead years capitalised", "OHYears", 30, "0")
r += 1

section(ws, r, "Forward market (nominal $/MWh ATC; PJM capacity $/MW-day)", 7)
r += 1
header(ws, r, ["Item", "2027", "2028", "2029"])
r += 1
put(ws, f"A{r}", "PJM West ATC forward")
for j, y in enumerate([2027, 2028, 2029]):
    put(ws, f"{L(2+j)}{r}", I.FORWARD_ATC["PJMW"][y], BLUE, USD1)
    name(f"FwdPJM{y-2000}", f"Inputs!${L(2+j)}${r}")
put(ws, f"G{r}", "Derived from on-peak Cal-27 $83.90 / Cal-28 $85.35 / Cal-29 $81.50 (8-Sep-26) x ~0.86", italic=True)
r += 1
put(ws, f"A{r}", "ERCOT North ATC forward")
for j, y in enumerate([2027, 2028, 2029]):
    put(ws, f"{L(2+j)}{r}", I.FORWARD_ATC["ERCOT"][y], BLUE, USD1)
    name(f"FwdERC{y-2000}", f"Inputs!${L(2+j)}${r}")
put(ws, f"G{r}", "Estimate: EIA STEO 2027 $47.39; forwards reported >$50", italic=True)
r += 1
put(ws, f"A{r}", "PJM capacity, calendar year (blend of delivery years)")
put(ws, f"B{r}", round(I.PJM_CAPACITY_CY[2027], 2), BLUE, USD2)
put(ws, f"C{r}", round(I.PJM_CAPACITY_CY[2028], 2), BLUE, USD2)
put(ws, f"D{r}", I.PJM_CAPACITY_KNOWN["2028/29"], BLUE, USD2, note="2029 column = 2028/29 delivery-year price; CY2029 blends 5/12 of it with 7/12 of long-run")
name("CapCY27", f"Inputs!$B${r}")
name("CapCY28", f"Inputs!$C${r}")
name("CapDY2829", f"Inputs!$D${r}")
put(ws, f"G{r}", "PJM BRA: 26/27 $329.17, 27/28 $333.44, 28/29 $325 (cap)", italic=True)
r += 2

section(ws, r, "Scenario table (long-run, 2026$)", 7)
r += 1
header(ws, r, ["Parameter", "Low", "Base", "High", "Extreme", "ACTIVE", "ACTIVE ERCOT"])
scen_hdr = r
r += 1
params = [("Implied heat rate PJM (MMBtu/MWh)", "ihr_pjm", "IHRPJM", USD2), ("Implied heat rate ERCOT", "ihr_ercot", "IHRERC", USD2),
          ("PJM capacity ($/MW-day)", "pjm_cap", "PJMCapScen", USD), ("NY/NE/MISO capacity ($/MW-day)", "other_cap", "OtherCap", USD),
          ("CAISO RA capacity ($/MW-day)", "west_cap", "WestCap", USD), ("Peaker scarcity capture (x ATC)", "scarcity_mult", "Scarcity", USD2),
          ("US data-center TWh 2030 (memo)", "us_dc_twh_2030", "DCTWh30", USD), ("US data-center TWh 2035 (memo)", "us_dc_twh_2035", "DCTWh35", USD)]
for lab, key, nm, fmt in params:
    put(ws, f"A{r}", lab)
    for j, s in enumerate(I.SCENARIO_ORDER):
        put(ws, f"{L(2+j)}{r}", I.SCENARIOS[s][key], BLUE, fmt)
    put(ws, f"F{r}", f"=INDEX(B{r}:E{r},MATCH(Scenario,$B${scen_hdr}:$E${scen_hdr},0))", fmt=fmt, bold=True)
    name(nm, f"Inputs!$F${r}")
    if key in ("ihr_ercot", "scarcity_mult"):
        put(ws, f"G{r}", f'=INDEX(B{r}:E{r},MATCH(IF(ERCOTScen="Same",Scenario,ERCOTScen),$B${scen_hdr}:$E${scen_hdr},0))', fmt=fmt, bold=True)
        name("IHRERC" if key == "ihr_ercot" else "ScarcityERC", f"Inputs!$G${r}")
    if key in ("ihr_pjm", "other_cap", "west_cap"):
        name(nm + "Base", f"Inputs!$C${r}")
    r += 1
put(ws, f"A{r}", "PJM capacity - active (override if set)")
put(ws, f"F{r}", "=IF(ISNUMBER(PJMCapOverride),PJMCapOverride,PJMCapScen)", fmt=USD, bold=True)
name("PJMCap", f"Inputs!$F${r}")
r += 2

section(ws, r, "Hubs", 7)
r += 1
header(ws, r, ["Hub", "Price ratio to PJM West", "Gas basis to HH ($/MMBtu)", "CCGT capacity factor", "CCGT capture (x ATC)", "Nuclear capture (x ATC)"])
r += 1
hub_row = {}
for h in HUBS:
    put(ws, f"A{r}", h)
    put(ws, f"B{r}", I.HUB_RATIO_TO_PJMW.get(h, 1.0) if h != "ERCOT" else None, BLUE, USD2)
    put(ws, f"C{r}", I.GAS_BASIS[h], BLUE, USD2)
    put(ws, f"D{r}", I.CCGT_CF.get(h, 0.45 if h in ("MISO", "NY") else 0.5), BLUE, PCT)
    put(ws, f"E{r}", I.CCGT_CAPTURE.get(h, 1.02 if h in ("NIHUB", "MISO") else 1.05), BLUE, USD2)
    put(ws, f"F{r}", I.NUCLEAR_CAPTURE.get(h, I.TECH["nuclear"]["capture"]), BLUE, USD2)
    hub_row[h] = r
    r += 1
put(ws, f"G{hub_row['ERCOT']}", "ERCOT prices use their own heat rate (no ratio)", italic=True)
r += 1

section(ws, r, "Technology defaults (asset sheets link here; overwrite per asset if you like)", 7)
r += 1
tcols = ["Tech", "Capacity factor", "Capture (x ATC)", "Heat rate", "Fuel price override (2026$/MMBtu)", "Variable O&M ($/MWh)",
         "Fuel ($/MWh)", "O&M ($/MWh)", "Capex ($/MWh)", "Per-MWh cost escalator", "Fixed O&M ($/kW-yr)", "Sust. capex ($/kW-yr)",
         "Capacity accreditation", "Merchant discount rate", "Contract discount rate", "Dispatch check (1/0)", "Mothball floor (1/0)",
         "45U eligible (1/0)", "Replacement cost ($/kW)"]
header(ws, r, tcols)
ws.row_dimensions[r].height = 45
r += 1
nuc = I.TECH["nuclear"]
tech_vals = {
    "nuclear": [nuc["cf"], nuc["capture"], 0, None, 0, 5.5, 22.3, 7.7, nuc["cost_esc"], 0, 0, I.ACCREDITATION["nuclear"], nuc["disc"], nuc["disc_contract"], 0, 0, 1, None],
    "ccgt": [None, None, I.TECH["ccgt"]["heat_rate"], None, I.TECH["ccgt"]["vom"], 0, 0, 0, 0, I.TECH["ccgt"]["fom_kw"], I.TECH["ccgt"]["capex_kw"], I.ACCREDITATION["ccgt"], I.TECH["ccgt"]["disc"], I.TECH["ccgt"]["disc"], 1, 1, 0, I.REPLACEMENT_COST_KW["ccgt"]],
    "peaker": [I.TECH["peaker"]["cf"], None, I.TECH["peaker"]["heat_rate"], None, I.TECH["peaker"]["vom"], 0, 0, 0, 0, I.TECH["peaker"]["fom_kw"], I.TECH["peaker"]["capex_kw"], I.ACCREDITATION["peaker"], I.TECH["peaker"]["disc"], I.TECH["peaker"]["disc"], 1, 1, 0, I.REPLACEMENT_COST_KW["peaker"]],
    "coal": [I.TECH["coal"]["cf"], I.TECH["coal"]["capture"], I.TECH["coal"]["heat_rate"], I.TECH["coal"]["fuel"], I.TECH["coal"]["vom"], 0, 0, 0, 0, I.TECH["coal"]["fom_kw"], I.TECH["coal"]["capex_kw"], I.ACCREDITATION["coal"], I.TECH["coal"]["disc"], I.TECH["coal"]["disc"], 1, 1, 0, I.REPLACEMENT_COST_KW["coal"]],
    "hydro": [I.TECH["hydro"]["cf"], I.TECH["hydro"]["capture"], 0, None, 0, 0, 0, 0, 0, I.TECH["hydro"]["fom_kw"], 0, I.ACCREDITATION["hydro"], I.TECH["hydro"]["disc"], I.TECH["hydro"]["disc"], 0, 0, 0, I.REPLACEMENT_COST_KW["hydro"]],
    "geothermal": [I.TECH["geothermal"]["cf"], 1.10, 0, None, 0, 0, I.TECH["geothermal"]["cost_mwh"], 0, 0.025, 0, 0, I.ACCREDITATION["geothermal"], I.TECH["geothermal"]["disc"], I.TECH["geothermal"]["disc"], 0, 0, 0, I.REPLACEMENT_COST_KW["geothermal"]],
    "renewable": [None] * 18 + [I.REPLACEMENT_COST_KW["renewable"]],
}
tech_vals["renewable"] = [None] * 17 + [I.REPLACEMENT_COST_KW["renewable"]]
tech_row = {}
fmts = [PCT, USD2, USD2, USD2, USD2, USD2, USD2, USD2, PCT, USD1, USD1, PCT, PCT, PCT, "0", "0", "0", USD]
for t, vals in tech_vals.items():
    put(ws, f"A{r}", t)
    for j, v in enumerate(vals):
        if v is not None:
            put(ws, f"{L(2+j)}{r}", v, BLUE, fmts[j])
    tech_row[t] = r
    r += 1
put(ws, f"A{r}", "Notes: nuclear per-MWh cost $35.5 = fuel 5.5 + O&M 22.3 + capex 7.7 (split estimated from NEI 'Nuclear Costs in Context', 2025 avg $36.46/MWh). "
    "CCGT CF/capture come from the hub table; peaker capture = scenario scarcity multiple; coal uses a fixed fuel price.", italic=True)
r += 2
TCOL = {c: L(2 + i) for i, c in enumerate(["cf", "capture", "hr", "fuelovr", "vom", "fuel_mwh", "om_mwh", "capex_mwh", "esc", "fom", "capex_kw",
                                           "accr", "disc_m", "disc_c", "dispatch", "mothball", "ptc", "rc"])}

section(ws, r, "Nuclear replacement-cost basis ($/kW new build)", 7)
r += 1
header(ws, r, ["#", "Basis", "$/kW"])
r += 1
first_nb = r
for i, (lab, v, src) in enumerate([("Gas-equivalent: new CCGT per accredited MW", round(I.NUCLEAR_FUNCTIONAL_KW), "$2,350 x 0.95/0.78"),
                                   ("Gas + clean premium (estimate)", 4500, "Gas-equivalent + PV of ~$15/MWh clean premium"),
                                   ("Nth-of-a-kind AP1000 (MIT, 4th plant)", 6200, "MIT CANES 2024"),
                                   ("EPRI next AP1000 (base)", I.REPLACEMENT_COST_KW["nuclear_like"], "EPRI $9,700-15,100; Vogtle actual ~$17,500-21,700")], 1):
    put(ws, f"A{r}", i, BLUE)
    put(ws, f"B{r}", lab)
    put(ws, f"C{r}", v, BLUE, USD)
    put(ws, f"G{r}", src, italic=True)
    r += 1
put(ws, f"B{r}", "Selected", bold=True)
put(ws, f"C{r}", f"=INDEX(C{first_nb}:C{r-1},NucBasis)", fmt=USD, bold=True)
name("NucRC", f"Inputs!$C${r}")
r += 2

section(ws, r, "Company inputs", 7)
r += 1
header(ws, r, ["Item", "CEG", "VST", "", "", "", "Source / note"])
r += 1
comp = {}


def cinp(label, key, vals, fmt, lever=False, src=""):
    global r
    put(ws, f"A{r}", label)
    for j, co in enumerate(("CEG", "VST")):
        put(ws, f"{L(2+j)}{r}", vals[j], BLUE, fmt, fill=YELLOW if lever else None)
        name(f"{co}_{key}", f"Inputs!${L(2+j)}${r}")
    if src:
        put(ws, f"G{r}", src, italic=True)
    r += 1


cinp("Share price ($)", "Price", [I.MARKET["CEG"]["price"], I.MARKET["VST"]["price"]], USD2, src="CEG 22-Sep-26; VST 25-Sep-26")
cinp("Shares outstanding (mm)", "Shares", [I.MARKET["CEG"]["shares_mm"], I.MARKET["VST"]["shares_mm"]], USD1, src="CEG 10-Q cover 31-Jul-26; VST ~336mm Aug-26")
cinp("Additional shares from pending deals (mm)", "ExtraShares", [0, 5], USD1, src="5mm VST shares to Cogentrix sellers")
cinp("Corporate overhead ($mm/yr, 2026$)", "OH", [I.CORP_OVERHEAD_MM["CEG"], I.CORP_OVERHEAD_MM["VST"]], USD, src="Estimate")
cinp("Retail/platform EBITDA ($mm/yr, 2026$)", "PlatEBITDA", [I.PLATFORM["CEG"]["ebitda_mm"], I.PLATFORM["VST"]["ebitda_mm"]], USD, src="VST Retail FY25 $1,463mm reported; CEG estimate")
cinp("Platform value multiple (x EBITDA)", "PlatMult", [6.0, 6.0], "0.0x")
cinp("Net debt bearing interest ($bn)", "NetDebt", [round(I.FINANCING["CEG"]["net_debt_bn"], 3), round(I.FINANCING["VST"]["net_debt_bn"], 3)], USD2, src="After pending-deal cash")
cinp("Cost of debt", "CoD", [I.FINANCING["CEG"]["cost_of_debt"], I.FINANCING["VST"]["cost_of_debt"]], PCT, src="Estimate")
cinp("Preferred stock ($bn)", "Pref", [0, I.FINANCING["VST"]["pref_bn"]], USD2)
cinp("Preferred dividend rate", "PrefRate", [0, I.FINANCING["VST"]["pref_rate"]], PCT)
cinp("Renewables cash yield (% of $/kW value)", "RenYield", [I.RENEWABLE_CASH_YIELD] * 2, PCT)
cinp("Revenue adjustment (all assets)", "RevAdj", [0, 0], PCT, lever=True, src="+10% = every revenue line x1.10")
cinp("Fuel & variable cost adjustment", "FuelAdj", [0, 0], PCT, lever=True, src="Moves gross margin")
cinp("Fixed O&M adjustment", "OMAdj", [0, 0], PCT, lever=True)
cinp("Capex adjustment (sustaining + one-off)", "CapexAdj", [0, 0], PCT, lever=True)
cinp("PPA price adjustment ($/MWh)", "PPAAdj", [0, 0], USD2, lever=True, src="Undisclosed contract prices - analyst estimates")
cinp("FCF haircut (years up to horizon)", "Haircut", [0, 0], PCT, lever=True, src="Model's 2027 FCF looks ~15-25% above guidance run-rate")
cinp("Other value ($mm): growth projects, site / data-center (Helix) options", "OtherVal", [0, 0], USD, lever=True,
     src="Not modelled by default. E.g. Talen sold its Cumulus data-center campus to Amazon for $650mm (2024)")
cinp("Price-target weight: NAV incl. platform", "W1", [1, 1], "0.00", lever=True)
cinp("Price-target weight: levered equity value (FCF to horizon + NAV at horizon, PV at equity rate)", "W2", [1, 1], "0.00", lever=True)
cinp("Price-target weight: replacement cost + FCF", "W3", [1, 1], "0.00", lever=True)
r += 1

section(ws, r, "Claims on the assets ($bn) - positive = liability, negative = asset", 7)
r += 1
header(ws, r, ["Item", "$bn", "", "", "", "", "Source / note"])
r += 1
for co, claims in (("CEG", I.CEG_CLAIMS), ("VST", I.VST_CLAIMS)):
    put(ws, f"A{r}", co, bold=True)
    r += 1
    first = r
    for lab, v in claims:
        put(ws, f"A{r}", lab)
        put(ws, f"B{r}", v, BLUE, USD2)
        r += 1
    put(ws, f"A{r}", f"{co} net claims", bold=True)
    put(ws, f"B{r}", f"=SUM(B{first}:B{r-1})", fmt=USD2, bold=True)
    name(f"{co}_Claims", f"Inputs!$B${r}")
    r += 2

# =====================================================================================
# PRICES
# =====================================================================================
pr = wb.create_sheet("Prices")
pr.column_dimensions["A"].width = 14
pr.column_dimensions["B"].width = 44
pr.column_dimensions["C"].width = 12
for c in YC:
    pr.column_dimensions[c].width = 9
put(pr, "A1", "PRICE DECK - formulas only (edit on Inputs)", bold=True)
put(pr, "A3", "Key", bold=True)
put(pr, "B3", "Year", bold=True)
for i, y in enumerate(YEARS):
    put(pr, f"{YC[i]}3", y, fmt=YRFMT, bold=True)
put(pr, "B4", "Inflation index (2026 = 1)")
put(pr, "A4", "INFL")
put(pr, "B5", "Henry Hub, real 2026$/MMBtu")
put(pr, "A5", "HH")
for c in YC:
    put(pr, f"{c}4", f"=(1+Infl)^({c}$3-2026)", fmt="0.000")
    put(pr, f"{c}5", f"=IF({c}$3=2027,HH27,IF({c}$3=2028,HH28,HHLR))+GasShift", fmt=USD2)
put(pr, "A7", "Long-run ATC (2026$/MWh) by hub", bold=True)
lr_row = {}
rr = 8
for h in HUBS:
    put(pr, f"A{rr}", f"LR|{h}")
    put(pr, f"B{rr}", f"{h} long-run ATC, 2026$")
    hr = hub_row[h]
    if h == "ERCOT":
        f = f"=IHRERC*(HHLR+GasShift+Inputs!$C${hr})+PriceShift"
    else:
        f = f"=IHRPJM*(HHLR+GasShift+Inputs!$C${hub_row['PJMW']})*Inputs!$B${hr}+PriceShift"
    put(pr, f"C{rr}", f, fmt=USD1, bold=True)
    lr_row[h] = rr
    rr += 1
put(pr, f"A{rr}", "LR|PJMW_BASE")
put(pr, f"B{rr}", "PJM West long-run ATC in Base scenario, no shifts (for $/kW-valued renewables)")
put(pr, f"C{rr}", f"=IHRPJMBase*(HHLR+Inputs!$C${hub_row['PJMW']})", fmt=USD1)
name("PJMWLRBase", f"Prices!$C${rr}")
name("PJMWLR", f"Prices!$C${lr_row['PJMW']}")
rr += 2
put(pr, f"A{rr}", "Annual deck (nominal)", bold=True)
rr += 1
for h in HUBS:
    hr = hub_row[h]
    # gas
    put(pr, f"A{rr}", f"GAS|{h}")
    put(pr, f"B{rr}", f"{h} gas, nominal $/MMBtu")
    for c in YC:
        put(pr, f"{c}{rr}", f"=({c}$5+Inputs!$C${hr})*{c}$4", fmt=USD2)
    rr += 1
    # ATC
    put(pr, f"A{rr}", f"ATC|{h}")
    put(pr, f"B{rr}", f"{h} ATC power, nominal $/MWh")
    for c in YC:
        if h == "ERCOT":
            fwd = f"IF({c}$3=2027,FwdERC27,FwdERC28)"
            f29 = "FwdERC29"
            ratio = "1"
        else:
            fwd = f"IF({c}$3=2027,FwdPJM27,FwdPJM28)"
            f29 = "FwdPJM29"
            ratio = f"Inputs!$B${hr}"
        put(pr, f"{c}{rr}", f"=IF({c}$3<=2028,{fwd}*{ratio},IF({c}$3=2029,0.5*{f29}*{ratio}+0.5*$C${lr_row[h]}*{c}$4,$C${lr_row[h]}*{c}$4))", fmt=USD1)
    rr += 1
    # capacity
    put(pr, f"A{rr}", f"CAP|{h}")
    put(pr, f"B{rr}", f"{h} capacity, nominal $/MW-day")
    for c in YC:
        if h in PJM_HUBS:
            f = f"=IF({c}$3=2027,CapCY27,IF({c}$3=2028,CapCY28,IF({c}$3=2029,5/12*CapDY2829+7/12*PJMCap*{c}$4,PJMCap*{c}$4)))"
        elif h == "ERCOT":
            f = "=0"
        elif h == "WEST":
            f = f"=IF({c}$3<=2028,WestCapBase,WestCap)*{c}$4"
        else:
            f = f"=IF({c}$3<=2028,OtherCapBase,OtherCap)*{c}$4"
        put(pr, f"{c}{rr}", f, fmt=USD1)
    rr += 1
pr.freeze_panes = "C4"


# =====================================================================================
# ASSET SHEETS
# =====================================================================================
ACOLS = ["Asset", "Tech", "Hub", "MW", "Start year", "End year", "Avg in-service year (COD)", "Method (Cashflow / $/kW)",
         "Capacity factor", "Capture (x ATC)", "Heat rate", "Fuel price override (2026$/MMBtu)", "Variable O&M ($/MWh)",
         "Fuel ($/MWh)", "O&M ($/MWh)", "Capex ($/MWh)", "Per-MWh cost escalator", "Fixed O&M ($/kW-yr)", "Sust. capex ($/kW-yr)",
         "Capacity accreditation", "ZEC ($/MWh)", "ZEC last year", "45U eligible", "Dispatch check", "Mothball floor",
         "Contract MW", "Contract start", "Contract end", "Contract ramp (yrs)", "PPA price ($/MWh, nominal flat)",
         "Merchant discount rate", "Contract discount rate", "One-off capex ($mm)", "One-off capex year", "$/kW value (if $/kW method)",
         "Replacement cost ($/kW)", "Remaining-life override", "Remaining-life fraction", "Depreciated replacement cost ($mm)",
         "ATC row", "Gas row", "Cap row", "Asset NAV ($mm)", "NAV $/kW"]
AC = {k: L(i + 1) for i, k in enumerate(["name", "tech", "hub", "mw", "start", "end", "cod", "method", "cf", "capture", "hr", "fuelovr", "vom",
                                        "fuel_mwh", "om_mwh", "capex_mwh", "esc", "fom", "capex_kw", "accr", "zec", "zec_end", "ptc", "dispatch",
                                        "mothball", "cmw", "cstart", "cend", "ramp", "ppa", "disc_m", "disc_c", "oneoff", "oneoff_yr", "usdkw",
                                        "rc", "lifeovr", "lifefrac", "drc", "atcrow", "gasrow", "caprow", "nav", "navkw"])}


def lookup(d, nm):
    for k, v in d.items():
        if nm.startswith(k) or k in nm:
            return v
    return None


def build_assets(co, assets):
    sh = wb.create_sheet(f"{co}_Assets")
    header(sh, 1, ACOLS)
    sh.row_dimensions[1].height = 60
    sh.column_dimensions["A"].width = 58
    for i in range(2, len(ACOLS) + 1):
        sh.column_dimensions[L(i)].width = 11
    for i, a in enumerate(assets):
        row = i + 2
        t = a["tech"]
        tr = tech_row.get(t)
        link = lambda key: f"=Inputs!${TCOL[key]}${tr}"
        put(sh, f"{AC['name']}{row}", a["name"], BLUE)
        put(sh, f"{AC['tech']}{row}", t, BLUE)
        put(sh, f"{AC['hub']}{row}", a["hub"], BLUE)
        put(sh, f"{AC['mw']}{row}", a["mw"], BLUE, USD)
        put(sh, f"{AC['cod']}{row}", lookup(I.ASSET_COD, a["name"]), BLUE, YRFMT)
        if t == "renewable":
            put(sh, f"{AC['method']}{row}", "$/kW", BLUE)
            put(sh, f"{AC['start']}{row}", 2027, BLUE, YRFMT)
            put(sh, f"{AC['end']}{row}", 2026, BLUE, YRFMT)
            put(sh, f"{AC['usdkw']}{row}", a["usd_per_kw"], BLUE, USD)
            for k in ("cf", "capture", "hr", "vom", "fuel_mwh", "om_mwh", "capex_mwh", "esc", "fom", "capex_kw", "accr", "zec", "ptc", "dispatch", "mothball", "cmw", "ppa", "oneoff"):
                put(sh, f"{AC[k]}{row}", 0, BLUE)
            for k in ("zec_end", "cstart", "cend", "oneoff_yr"):
                put(sh, f"{AC[k]}{row}", 2026, BLUE, YRFMT)
            put(sh, f"{AC['ramp']}{row}", 1, BLUE)
            put(sh, f"{AC['disc_m']}{row}", 0.08, BLUE, PCT)
            put(sh, f"{AC['disc_c']}{row}", 0.08, BLUE, PCT)
        else:
            put(sh, f"{AC['method']}{row}", "Cashflow", BLUE)
            put(sh, f"{AC['start']}{row}", a["start"], BLUE, YRFMT)
            put(sh, f"{AC['end']}{row}", a["end"], BLUE, YRFMT)
            hr_ = hub_row[a["hub"]]
            if t == "ccgt":
                put(sh, f"{AC['cf']}{row}", f"=Inputs!$D${hr_}", GREEN, PCT)
                put(sh, f"{AC['capture']}{row}", f"=Inputs!$E${hr_}", GREEN, USD2)
            elif t == "peaker":
                put(sh, f"{AC['cf']}{row}", link("cf"), GREEN, PCT)
                put(sh, f"{AC['capture']}{row}", "=ScarcityERC" if a["hub"] == "ERCOT" else "=Scarcity", GREEN, USD2)
            elif t == "nuclear":
                put(sh, f"{AC['cf']}{row}", link("cf"), GREEN, PCT)
                put(sh, f"{AC['capture']}{row}", f"=Inputs!$F${hr_}", GREEN, USD2)
            else:
                put(sh, f"{AC['cf']}{row}", link("cf"), GREEN, PCT)
                put(sh, f"{AC['capture']}{row}", link("capture"), GREEN, USD2)
            if a.get("heat_rate"):
                put(sh, f"{AC['hr']}{row}", a["heat_rate"], BLUE, USD2)
            else:
                put(sh, f"{AC['hr']}{row}", link("hr"), GREEN, USD2)
            if t == "coal":
                put(sh, f"{AC['fuelovr']}{row}", link("fuelovr"), GREEN, USD2)
            for k, fmt in (("vom", USD2), ("fuel_mwh", USD2), ("om_mwh", USD2), ("capex_mwh", USD2), ("esc", PCT), ("fom", USD1),
                           ("capex_kw", USD1), ("accr", PCT), ("ptc", "0"), ("dispatch", "0"), ("mothball", "0"), ("disc_m", PCT), ("disc_c", PCT)):
                put(sh, f"{AC[k]}{row}", link(k), GREEN, fmt)
            z = a.get("zec")
            put(sh, f"{AC['zec']}{row}", z["price"] if z else 0, BLUE, USD2)
            put(sh, f"{AC['zec_end']}{row}", z["end"] if z else 2026, BLUE, YRFMT)
            c = (a.get("contracts") or [None])[0]
            if c:
                full = c.get("mw_full", c["mw"])
                ramp = {"Nuclear - Comanche": 5, "Nuclear - Perry": 2}.get(next((k for k in ("Nuclear - Comanche", "Nuclear - Perry") if a["name"].startswith(k)), ""), 1)
                put(sh, f"{AC['cmw']}{row}", full, BLUE, USD)
                put(sh, f"{AC['cstart']}{row}", c["start"], BLUE, YRFMT)
                put(sh, f"{AC['cend']}{row}", c["end"], BLUE, YRFMT)
                put(sh, f"{AC['ramp']}{row}", ramp, BLUE, "0")
                put(sh, f"{AC['ppa']}{row}", c["price"], BLUE, USD2, note="Undisclosed - analyst estimate")
            else:
                put(sh, f"{AC['cmw']}{row}", 0, BLUE, USD)
                put(sh, f"{AC['cstart']}{row}", 2026, BLUE, YRFMT)
                put(sh, f"{AC['cend']}{row}", 2026, BLUE, YRFMT)
                put(sh, f"{AC['ramp']}{row}", 1, BLUE, "0")
                put(sh, f"{AC['ppa']}{row}", 0, BLUE, USD2)
            capex = a.get("capex", {})
            put(sh, f"{AC['oneoff']}{row}", sum(capex.values()) * 1000 if capex else 0, BLUE, USD)
            put(sh, f"{AC['oneoff_yr']}{row}", list(capex)[0] if capex else 2026, BLUE, YRFMT)
        # replacement cost
        if t == "nuclear":
            put(sh, f"{AC['rc']}{row}", "=NucRC", GREEN, USD)
        else:
            put(sh, f"{AC['rc']}{row}", link("rc"), GREEN, USD)
        ov = lookup(I.LIFE_FRAC_OVERRIDE, a["name"])
        if ov is not None:
            put(sh, f"{AC['lifeovr']}{row}", ov, BLUE, PCT)
        put(sh, f"{AC['lifefrac']}{row}",
            f"=IF(ISNUMBER({AC['lifeovr']}{row}),{AC['lifeovr']}{row},MAX(0,MIN(1,({AC['end']}{row}+1-ValDate)/({AC['end']}{row}+1-{AC['cod']}{row}))))", fmt=PCT)
        put(sh, f"{AC['drc']}{row}", f"={AC['mw']}{row}*{AC['rc']}{row}*{AC['lifefrac']}{row}/1000", fmt=USD)
        put(sh, f"{AC['atcrow']}{row}", f'=MATCH("ATC|"&{AC["hub"]}{row},Prices!$A:$A,0)', fmt="0")
        put(sh, f"{AC['gasrow']}{row}", f'=MATCH("GAS|"&{AC["hub"]}{row},Prices!$A:$A,0)', fmt="0")
        put(sh, f"{AC['caprow']}{row}", f'=MATCH("CAP|"&{AC["hub"]}{row},Prices!$A:$A,0)', fmt="0")
    last = len(assets) + 1
    tot = last + 1
    put(sh, f"A{tot}", "Total", bold=True)
    put(sh, f"{AC['mw']}{tot}", f"=SUM({AC['mw']}2:{AC['mw']}{last})", fmt=USD, bold=True)
    put(sh, f"{AC['drc']}{tot}", f"=SUM({AC['drc']}2:{AC['drc']}{last})", fmt=USD, bold=True)
    put(sh, f"{AC['nav']}{tot}", f"=SUM({AC['nav']}2:{AC['nav']}{last})", fmt=USD, bold=True)
    sh.freeze_panes = "B2"
    return sh, last


# =====================================================================================
# MODEL SHEETS
# =====================================================================================
LINES = [  # key, label, fmt
    ("op", "Operating (1/0)", "0"), ("cmw", "Contracted MW", USD), ("gen_m", "Merchant generation (MWh)", USD),
    ("gen_c", "Contracted generation (MWh)", USD), ("fuelp", "Fuel price, delivered ($/MMBtu)", USD2), ("spark", "Spread vs fuel ($/MWh)", USD2),
    ("run", "Runs (1/0)", "0"), ("energy", "Merchant energy revenue ($mm)", USD1), ("capacity", "Capacity revenue ($mm)", USD1),
    ("ppa", "Long-term PPA revenue ($mm)", USD1), ("zec", "ZEC revenue ($mm)", USD1), ("revenue", "Total revenue ($mm)", USD1),
    ("fuelvar", "Fuel & variable cost ($mm)", USD1), ("gm", "Gross margin ($mm)", USD1), ("fom", "Fixed O&M ($mm)", USD1),
    ("ebitda", "Asset EBITDA ($mm)", USD1), ("capex", "Capex - sustaining + one-off ($mm)", USD1), ("pretax", "Pre-tax cash flow ($mm)", USD1),
    ("ptc", "45U credit - untaxed ($mm)", USD1), ("con_cf", "  of which contracted pre-tax ($mm)", USD1), ("merch_cf", "  of which merchant pre-tax ($mm)", USD1),
    ("at_m", "After-tax merchant CF incl. 45U ($mm)", USD1), ("at_c", "After-tax contracted CF ($mm)", USD1),
    ("df_m", "Discount factor - merchant", "0.000"), ("df_c", "Discount factor - contracted", "0.000"),
    ("pv_m", "PV merchant ($mm)", USD1), ("pv_c", "PV contracted ($mm)", USD1), ("fixedval", "$/kW-method value ($mm)", USD1),
]


def build_model(co, assets, ash, alast):
    ms = wb.create_sheet(f"{co}_Model")
    ms.column_dimensions["A"].width = 46
    ms.column_dimensions["B"].width = 10
    for c in "CDE":
        ms.column_dimensions[c].width = 13
    ms.column_dimensions["F"].width = 2
    for c in YC:
        ms.column_dimensions[c].width = 10
    A = f"{co}_Assets"
    put(ms, "A1", f"{co} - annual asset cash flows and consolidated FCF ($mm, nominal). Formulas only.", bold=True)
    put(ms, "A3", "Line", bold=True)
    put(ms, "B3", "Key", bold=True)
    put(ms, "C3", "PV / total", bold=True)
    put(ms, "D3", "PV to horizon", bold=True)
    put(ms, "E3", "Value at horizon", bold=True)
    for i, y in enumerate(YEARS):
        put(ms, f"{YC[i]}3", y, fmt=YRFMT, bold=True)
    put(ms, "A4", "Inflation index")
    for c in YC:
        put(ms, f"{c}4", f"=Prices!{c}$4", GREEN, "0.000")

    # consolidated block rows are placed at top; asset blocks start below
    CONS = [("c_energy", "Merchant energy revenue", "energy"), ("c_capacity", "Capacity revenue", "capacity"), ("c_ppa", "Long-term PPA revenue", "ppa"),
            ("c_zec", "ZEC revenue", "zec"), ("c_rev", "TOTAL REVENUE", "revenue"), ("c_fuel", "Fuel & variable cost", "fuelvar"),
            ("c_gm", "GROSS MARGIN", "gm"), ("c_gmpct", "Gross margin %", None), ("c_fom", "Fixed O&M", "fom"), ("c_ebitda", "ASSET EBITDA", "ebitda"),
            ("c_ebitdapct", "EBITDA margin %", None), ("c_capex", "Capex (sustaining + one-off)", "capex"), ("c_pretax", "Asset pre-tax cash flow (after mothball floor)", "pretax"),
            ("c_ptc", "45U credit", "ptc"), ("c_ren", "Renewables cash (value x yield)", None), ("c_assetcash", "ASSET CASH MARGIN", None),
            ("c_plat", "+ Retail/platform EBITDA", None), ("c_oh", "- Corporate overhead", None), ("c_int", "- Interest", None),
            ("c_taxable", "Taxable income", None), ("c_tax", "- Cash tax", None), ("c_pref", "- Preferred dividends", None),
            ("c_fcf", "EQUITY FCF (before haircut)", None), ("c_fcfadj", "EQUITY FCF (after haircut)", None),
            ("c_ohpv", "Overhead PV (NAV, after tax)", None), ("c_ohfwd", "Overhead PV at horizon (after tax)", None), ("c_dfeq", "Equity discount factor", None)]
    crow = {}
    rr = 6
    put(ms, "A5", "CONSOLIDATED", bold=True)
    for key, lab, _ in CONS:
        crow[key] = rr
        rr += 1
    blk_start = rr + 2
    nlines = len(LINES)
    blk_end = blk_start + len(assets) * (nlines + 2) - 1
    BR = f"$B${blk_start}:$B${blk_end}"

    # ---- asset blocks
    for i, a in enumerate(assets):
        ar = i + 2
        p = lambda k: f"'{A}'!${AC[k]}${ar}"
        b0 = blk_start + i * (nlines + 2)
        put(ms, f"A{b0}", f"=" + p("name"), GREEN, bold=True)
        ms[f"A{b0}"].fill = SEC
        row = {k: b0 + 1 + j for j, (k, _, _) in enumerate(LINES)}
        for k, lab, fmt in LINES:
            put(ms, f"A{row[k]}", lab)
            put(ms, f"B{row[k]}", k)
        for c in YC:
            y = f"{c}$3"
            infl = f"{c}$4"
            atc = f"INDEX(Prices!{c}:{c},{p('atcrow')})"
            gas = f"INDEX(Prices!{c}:{c},{p('gasrow')})"
            cap = f"INDEX(Prices!{c}:{c},{p('caprow')})"
            esc = f"(1+{p('esc')})^({y}-2026)"
            R = lambda k: f"{c}{row[k]}"
            f = {
                "op": f'=IF(AND({y}>={p("start")},{y}<={p("end")},{p("method")}="Cashflow"),1,0)',
                "cmw": f"=IF(AND({y}>={p('cstart')},{y}<={p('cend')}),MIN({p('mw')},{p('cmw')}*MIN(1,({y}-{p('cstart')}+1)/MAX(1,{p('ramp')}))),0)*{R('op')}",
                "gen_m": f"=({p('mw')}-{R('cmw')})*8760*{p('cf')}*{R('op')}",
                "gen_c": f"={R('cmw')}*8760*{p('cf')}*{R('op')}",
                "fuelp": f"=IF(ISNUMBER({p('fuelovr')}),{p('fuelovr')}*{infl},{gas}+FuelAdder*{infl})",
                "spark": f"={atc}*{p('capture')}-{p('hr')}*{R('fuelp')}-{p('vom')}*{infl}",
                "run": f"=IF({p('dispatch')}=1,IF({R('spark')}>0,1,0),1)",
                "energy": f"={R('gen_m')}*{atc}*{p('capture')}*{R('run')}/1000000*(1+{co}_RevAdj)",
                "capacity": f"=({p('mw')}-{R('cmw')})*{cap}*365*{p('accr')}*{R('op')}/1000000*(1+{co}_RevAdj)",
                "ppa": f"={R('gen_c')}*({p('ppa')}+{co}_PPAAdj)/1000000*(1+{co}_RevAdj)",
                "zec": f"=IF({y}<={p('zec_end')},{R('gen_m')}*{p('zec')},0)/1000000*(1+{co}_RevAdj)",
                "revenue": f"={R('energy')}+{R('capacity')}+{R('ppa')}+{R('zec')}",
                "fuelvar": f"=(({p('hr')}*{R('fuelp')}+{p('vom')}*{infl})*{R('gen_m')}*{R('run')}+{p('fuel_mwh')}*{esc}*({R('gen_m')}+{R('gen_c')}))/1000000*(1+{co}_FuelAdj)",
                "gm": f"={R('revenue')}-{R('fuelvar')}",
                "fom": f"=({p('fom')}*{p('mw')}*1000*{infl}*{R('op')}+{p('om_mwh')}*{esc}*({R('gen_m')}+{R('gen_c')}))/1000000*(1+{co}_OMAdj)",
                "ebitda": f"={R('gm')}-{R('fom')}",
                "capex": f"=(({p('capex_kw')}*{p('mw')}*1000*{infl}*{R('op')}+{p('capex_mwh')}*{esc}*({R('gen_m')}+{R('gen_c')}))/1000000+IF({y}={p('oneoff_yr')},{p('oneoff')},0))*(1+{co}_CapexAdj)",
                "pretax": f"=IF({p('mothball')}=1,MAX(0,{R('ebitda')}-{R('capex')}),{R('ebitda')}-{R('capex')})",
                "ptc": f"=IF(AND({p('ptc')}=1,{y}<=PTCLast,{R('gen_m')}>0),{R('gen_m')}*MAX(0,PTCCredit*(1+Infl)^({y}-PTCBase)-0.8*MAX(0,({R('energy')}+{R('capacity')}+{R('zec')})*1000000/MAX(1,{R('gen_m')})-PTCThresh*(1+Infl)^({y}-PTCBase)))/1000000,0)",
                "con_cf": f"={R('ppa')}-{R('gen_c')}*({p('fuel_mwh')}*(1+{co}_FuelAdj)+{p('om_mwh')}*(1+{co}_OMAdj)+{p('capex_mwh')}*(1+{co}_CapexAdj))*{esc}/1000000",
                "merch_cf": f"={R('pretax')}-{R('con_cf')}",
                "at_m": f"={R('merch_cf')}*(1-TaxRate)+{R('ptc')}",
                "at_c": f"={R('con_cf')}*(1-TaxRate)",
                "df_m": f"=(1+{p('disc_m')})^-({y}+0.5-ValDate)",
                "df_c": f"=(1+{p('disc_c')})^-({y}+0.5-ValDate)",
                "pv_m": f"={R('at_m')}*{R('df_m')}",
                "pv_c": f"={R('at_c')}*{R('df_c')}",
                "fixedval": "=0",
            }
            for k, _, fmt in LINES:
                put(ms, f"{c}{row[k]}", f[k], fmt=fmt)
        # summary columns
        for k in ("energy", "capacity", "ppa", "zec", "revenue", "fuelvar", "gm", "fom", "ebitda", "capex", "pretax", "ptc"):
            put(ms, f"C{row[k]}", f"=SUM({YL}{row[k]}:{YR}{row[k]})", fmt=USD)
        for k, dk in (("pv_m", "disc_m"), ("pv_c", "disc_c")):
            put(ms, f"C{row[k]}", f"=SUM({YL}{row[k]}:{YR}{row[k]})", fmt=USD1, bold=True)
            put(ms, f"D{row[k]}", f'=SUMIF(${YL}$3:${YR}$3,"<="&Horizon,{YL}{row[k]}:{YR}{row[k]})', fmt=USD1)
            put(ms, f"E{row[k]}", f"=(C{row[k]}-D{row[k]})*(1+{p(dk)})^(Horizon+1-ValDate)", fmt=USD1)
        fv = f"=IF({p('method')}=\"$/kW\",{p('mw')}*{p('usdkw')}*(1+0.4*(PJMWLR/PJMWLRBase-1))/1000,0)"
        put(ms, f"C{row['fixedval']}", fv, fmt=USD1, bold=True)
        put(ms, f"E{row['fixedval']}", f"=C{row['fixedval']}", fmt=USD1)
        # link back to asset sheet
        ash[f"{AC['nav']}{ar}"] = f"='{co}_Model'!C{row['pv_m']}+'{co}_Model'!C{row['pv_c']}+'{co}_Model'!C{row['fixedval']}"
        ash[f"{AC['nav']}{ar}"].font = F(GREEN)
        ash[f"{AC['nav']}{ar}"].number_format = USD
        ash[f"{AC['navkw']}{ar}"] = f"=IF({AC['mw']}{ar}=0,0,{AC['nav']}{ar}/{AC['mw']}{ar}*1000)"
        ash[f"{AC['navkw']}{ar}"].font = F()
        ash[f"{AC['navkw']}{ar}"].number_format = USD

    # ---- consolidated rows
    for key, lab, src in CONS:
        rrow = crow[key]
        put(ms, f"A{rrow}", lab, bold=lab.isupper())
        put(ms, f"B{rrow}", key)
        for c in YC:
            y = f"{c}$3"
            Cc = lambda k: f"{c}{crow[k]}"
            if src:
                f = f'=SUMIF({BR},"{src}",{c}${blk_start}:{c}${blk_end})'
            else:
                f = {
                    "c_gmpct": f"=IF({Cc('c_rev')}=0,0,{Cc('c_gm')}/{Cc('c_rev')})",
                    "c_ebitdapct": f"=IF({Cc('c_rev')}=0,0,{Cc('c_ebitda')}/{Cc('c_rev')})",
                    "c_ren": f'=SUMIF({BR},"fixedval",$C${blk_start}:$C${blk_end})*{co}_RenYield',
                    "c_assetcash": f"={Cc('c_pretax')}+{Cc('c_ptc')}+{Cc('c_ren')}",
                    "c_plat": f"={co}_PlatEBITDA*{c}$4",
                    "c_oh": f"={co}_OH*{c}$4",
                    "c_int": f"={co}_NetDebt*1000*{co}_CoD",
                    "c_taxable": f"={Cc('c_assetcash')}-{Cc('c_ptc')}+{Cc('c_plat')}-{Cc('c_oh')}-{Cc('c_int')}",
                    "c_tax": f"=TaxRate*MAX(0,{Cc('c_taxable')})",
                    "c_pref": f"={co}_Pref*1000*{co}_PrefRate",
                    "c_fcf": f"={Cc('c_assetcash')}+{Cc('c_plat')}-{Cc('c_oh')}-{Cc('c_int')}-{Cc('c_tax')}-{Cc('c_pref')}",
                    "c_fcfadj": f"={Cc('c_fcf')}*(1-IF({y}<=Horizon,{co}_Haircut,0))",
                    "c_ohpv": f"=IF({y}<=2026+OHYears,{co}_OH*{c}$4*(1-TaxRate)*(1+OHDisc)^-({y}+0.5-ValDate),0)",
                    "c_ohfwd": f"=IF(AND({y}>Horizon,{y}<=Horizon+OHYears),{co}_OH*{c}$4*(1-TaxRate)*(1+OHDisc)^-({y}+0.5-(Horizon+1)),0)",
                    "c_dfeq": f"=(1+EquityDisc)^-({y}+0.5-ValDate)",
                }[key]
            fmt = PCT if key.endswith("pct") else ("0.000" if key == "c_dfeq" else USD)
            put(ms, f"{c}{rrow}", f, fmt=fmt, bold=lab.isupper())
        if key not in ("c_gmpct", "c_ebitdapct", "c_dfeq"):
            put(ms, f"C{rrow}", f"=SUM({YL}{rrow}:{YR}{rrow})", fmt=USD, bold=lab.isupper())
    for k in ("c_rev", "c_gm", "c_ebitda", "c_assetcash", "c_fcf", "c_fcfadj"):
        for col in range(1, Y0 + len(YEARS)):
            ms.cell(row=crow[k], column=col).border = THIN
    ms.freeze_panes = f"{YL}4"
    return ms, crow, blk_start, blk_end


built = {}
for co, assets in (("CEG", I.CEG_ASSETS), ("VST", I.VST_ASSETS)):
    ash, alast = build_assets(co, assets)
    ms, crow, bs, be = build_model(co, assets, ash, alast)
    built[co] = (ash, alast, ms, crow, bs, be)

# =====================================================================================
# SUMMARY
# =====================================================================================
sm = wb.create_sheet("Summary", 1)
sm.column_dimensions["A"].width = 64
for c in "BCDEFGH":
    sm.column_dimensions[c].width = 15
put(sm, "A1", "SUMMARY - price target, NAV, FCF and replacement cost (recalculates from Inputs)", bold=True)
put(sm, "A2", '="Active scenario: "&Scenario&"   |   price shift $"&TEXT(PriceShift,"0.0")&"/MWh   |   gas shift $"&TEXT(GasShift,"0.00")&"/MMBtu   |   nuclear replacement $"&TEXT(NucRC,"#,##0")&"/kW"', italic=True)
header(sm, 4, ["Item", "CEG", "VST", "", "Notes"])
r = 5
srow = {}


def srow_add(key, label, fceg, fvst, fmt, bold=False, note=""):
    global r
    put(sm, f"A{r}", label, bold=bold)
    put(sm, f"B{r}", fceg, fmt=fmt, bold=bold)
    put(sm, f"C{r}", fvst, fmt=fmt, bold=bold)
    if note:
        put(sm, f"E{r}", note, italic=True)
    srow[key] = r
    r += 1


def both(tmpl):
    return tmpl.replace("{co}", "CEG"), tmpl.replace("{co}", "VST")


def sec(t):
    global r
    r += 1
    for col in range(1, 6):
        sm.cell(row=r, column=col).fill = SEC
    put(sm, f"A{r}", t, bold=True)
    r += 1


def crow_ref(co, k, col="C"):
    ms, crow = built[co][2], built[co][3]
    return f"'{co}_Model'!${col}${crow[k]}"


def blk_sum(co, key, col):
    bs, be = built[co][4], built[co][5]
    return f"SUMIF('{co}_Model'!$B${bs}:$B${be},\"{key}\",'{co}_Model'!${col}${bs}:${col}${be})"


sec("Market")
srow_add("price", "Share price ($)", "=CEG_Price", "=VST_Price", USD2)
srow_add("shares", "Shares incl. pending issuance (mm)", "=CEG_Shares+CEG_ExtraShares", "=VST_Shares+VST_ExtraShares", USD1)
srow_add("mcap", "Market cap ($mm)", f"=B{r-2}*B{r-1}", f"=C{r-2}*C{r-1}", USD)
srow_add("claims", "Net debt, preferred & other claims ($mm)", "=CEG_Claims*1000", "=VST_Claims*1000", USD)
srow_add("ev", "Enterprise value ($mm)", f"=B{srow['mcap']}+B{srow['claims']}", f"=C{srow['mcap']}+C{srow['claims']}", USD)
srow_add("mw", "Owned MW (pro forma)", f"='CEG_Assets'!{AC['mw']}{built['CEG'][1]+1}", f"='VST_Assets'!{AC['mw']}{built['VST'][1]+1}", USD)
srow_add("evkw", "EV per kW ($)", f"=B{srow['ev']}/B{srow['mw']}*1000", f"=C{srow['ev']}/C{srow['mw']}*1000", USD)

sec("1. Asset NAV (sum of asset cash-flow values)")
srow_add("gav", "Gross asset value ($mm)", "=" + blk_sum("CEG", "pv_m", "C") + "+" + blk_sum("CEG", "pv_c", "C") + "+" + blk_sum("CEG", "fixedval", "C"),
         "=" + blk_sum("VST", "pv_m", "C") + "+" + blk_sum("VST", "pv_c", "C") + "+" + blk_sum("VST", "fixedval", "C"), USD, True)
srow_add("gavkw", "  GAV per kW ($)", f"=B{srow['gav']}/B{srow['mw']}*1000", f"=C{srow['gav']}/C{srow['mw']}*1000", USD)
srow_add("oh", "Less: capitalised corporate overhead ($mm)", "=" + crow_ref("CEG", "c_ohpv"), "=" + crow_ref("VST", "c_ohpv"), USD)
srow_add("eqnav", "Equity NAV - generation only ($mm)", f"=B{srow['gav']}-B{srow['oh']}-B{srow['claims']}", f"=C{srow['gav']}-C{srow['oh']}-C{srow['claims']}", USD, True)
srow_add("navps", "NAV per share - generation only ($)", f"=B{srow['eqnav']}/B{srow['shares']}", f"=C{srow['eqnav']}/C{srow['shares']}", USD2, True)
srow_add("plat", "Add: retail/platform value ($mm)", "=CEG_PlatEBITDA*CEG_PlatMult", "=VST_PlatEBITDA*VST_PlatMult", USD, note="Earnings-based; shown separately")
srow_add("other", "Add: other value - growth / site options ($mm)", "=CEG_OtherVal", "=VST_OtherVal", USD, note="Input on Inputs; 0 by default")
srow_add("navpsp", "NAV per share incl. platform & other ($)", f"=(B{srow['eqnav']}+B{srow['plat']}+B{srow['other']})/B{srow['shares']}", f"=(C{srow['eqnav']}+C{srow['plat']}+C{srow['other']})/C{srow['shares']}", USD2, True)

sec("2. Forward value: FCF generated to horizon + NAV at horizon")
srow_add("cumfcf", "Cumulative equity FCF to horizon ($mm, after haircut)",
         "=SUMIF('CEG_Model'!$" + YL + "$3:$" + YR + "$3,\"<=\"&Horizon,'CEG_Model'!$" + YL + "$" + str(built["CEG"][3]["c_fcfadj"]) + ":$" + YR + "$" + str(built["CEG"][3]["c_fcfadj"]) + ")",
         "=SUMIF('VST_Model'!$" + YL + "$3:$" + YR + "$3,\"<=\"&Horizon,'VST_Model'!$" + YL + "$" + str(built["VST"][3]["c_fcfadj"]) + ":$" + YR + "$" + str(built["VST"][3]["c_fcfadj"]) + ")", USD)
srow_add("fwdgav", "Gross asset value at horizon ($mm)", "=" + blk_sum("CEG", "pv_m", "E") + "+" + blk_sum("CEG", "pv_c", "E") + "+" + blk_sum("CEG", "fixedval", "E"),
         "=" + blk_sum("VST", "pv_m", "E") + "+" + blk_sum("VST", "pv_c", "E") + "+" + blk_sum("VST", "fixedval", "E"), USD)
srow_add("fwdoh", "Less: overhead at horizon ($mm)", "=" + crow_ref("CEG", "c_ohfwd"), "=" + crow_ref("VST", "c_ohfwd"), USD)
srow_add("fwdplat", "Add: platform value at horizon ($mm)", f"=B{srow['plat']}*(1+Infl)^(Horizon-2026)", f"=C{srow['plat']}*(1+Infl)^(Horizon-2026)", USD)
srow_add("fwdeq", "Equity NAV at horizon incl. other value ($mm)", f"=B{srow['fwdgav']}-B{srow['fwdoh']}-B{srow['claims']}+B{srow['fwdplat']}+B{srow['other']}",
         f"=C{srow['fwdgav']}-C{srow['fwdoh']}-C{srow['claims']}+C{srow['fwdplat']}+C{srow['other']}", USD)
srow_add("fwdps", "Value per share at horizon = FCF + NAV ($)", f"=(B{srow['cumfcf']}+B{srow['fwdeq']})/B{srow['shares']}",
         f"=(C{srow['cumfcf']}+C{srow['fwdeq']})/C{srow['shares']}", USD2, True)
srow_add("yrs", "Years to horizon", "=Horizon+1-ValDate", "=Horizon+1-ValDate", "0.00")
srow_add("irr", "Implied annual return from current price", f"=IF(B{srow['fwdps']}>0,(B{srow['fwdps']}/B{srow['price']})^(1/B{srow['yrs']})-1,-1)",
         f"=IF(C{srow['fwdps']}>0,(C{srow['fwdps']}/C{srow['price']})^(1/C{srow['yrs']})-1,-1)", PCT, True)
def lev(co, col):
    cr = built[co][3]
    yrs_rng = f"'{co}_Model'!${YL}$3:${YR}$3"
    fcf = f"'{co}_Model'!${YL}${cr['c_fcfadj']}:${YR}${cr['c_fcfadj']}"
    dfe = f"'{co}_Model'!${YL}${cr['c_dfeq']}:${YR}${cr['c_dfeq']}"
    return (f"=(SUMPRODUCT(({yrs_rng}<=Horizon)*{fcf}*{dfe})+{col}{srow['fwdeq']}*(1+EquityDisc)^-{col}{srow['yrs']})/{col}{srow['shares']}")


srow_add("fwdpv", "LEVERED EQUITY VALUE today: PV of FCF to horizon + PV of NAV at horizon ($/share)", lev("CEG", "B"), lev("VST", "C"), USD2, True,
         note="Discounted at the equity rate; credits cheap debt and the interest tax shield")
srow_add("levgap", "  of which cheap debt & interest tax shield vs asset NAV ($/share)", f"=B{srow['fwdpv']}-B{srow['navpsp']}", f"=C{srow['fwdpv']}-C{srow['navpsp']}", USD2)
srow_add("entry", "Entry price for target return (fat pitch, $)", f"=MAX(0,B{srow['fwdps']})/(1+TargetReturn)^B{srow['yrs']}",
         f"=MAX(0,C{srow['fwdps']})/(1+TargetReturn)^C{srow['yrs']}", USD2, True)

sec("3. Replacement cost + FCF")
srow_add("drc", "Depreciated replacement cost of fleet ($mm)", f"='CEG_Assets'!{AC['drc']}{built['CEG'][1]+1}", f"='VST_Assets'!{AC['drc']}{built['VST'][1]+1}", USD)
srow_add("eqdrc", "Less claims, plus other value = equity replacement value ($mm)", f"=B{srow['drc']}-B{srow['claims']}+B{srow['other']}", f"=C{srow['drc']}-C{srow['claims']}+C{srow['other']}", USD)
srow_add("rcfcf", "PV of equity FCF over replacement lead time ($mm)",
         "=SUMPRODUCT(('CEG_Model'!$" + YL + "$3:$" + YR + "$3<=2026+RCYears)*'CEG_Model'!$" + YL + f"${built['CEG'][3]['c_fcfadj']}:${YR}${built['CEG'][3]['c_fcfadj']}*'CEG_Model'!${YL}${built['CEG'][3]['c_dfeq']}:${YR}${built['CEG'][3]['c_dfeq']})",
         "=SUMPRODUCT(('VST_Model'!$" + YL + "$3:$" + YR + "$3<=2026+RCYears)*'VST_Model'!$" + YL + f"${built['VST'][3]['c_fcfadj']}:${YR}${built['VST'][3]['c_fcfadj']}*'VST_Model'!${YL}${built['VST'][3]['c_dfeq']}:${YR}${built['VST'][3]['c_dfeq']})", USD)
srow_add("rcps", "Replacement cost + FCF per share ($)", f"=(B{srow['eqdrc']}+B{srow['rcfcf']})/B{srow['shares']}", f"=(C{srow['eqdrc']}+C{srow['rcfcf']})/C{srow['shares']}", USD2, True)
srow_add("q", "Tobin's q (GAV / replacement cost)", f"=B{srow['gav']}/B{srow['drc']}", f"=C{srow['gav']}/C{srow['drc']}", MULT)
srow_add("evdrc", "Market EV / replacement cost", f"=B{srow['ev']}/B{srow['drc']}", f"=C{srow['ev']}/C{srow['drc']}", MULT)

sec("4. PRICE TARGET (weighted blend - weights on Inputs)")
srow_add("pt", "PRICE TARGET ($)",
         f"=(CEG_W1*B{srow['navpsp']}+CEG_W2*B{srow['fwdpv']}+CEG_W3*B{srow['rcps']})/MAX(0.0001,CEG_W1+CEG_W2+CEG_W3)",
         f"=(VST_W1*C{srow['navpsp']}+VST_W2*C{srow['fwdpv']}+VST_W3*C{srow['rcps']})/MAX(0.0001,VST_W1+VST_W2+VST_W3)", USD2, True)
srow_add("updown", "Upside / (downside) vs current price", f"=B{srow['pt']}/B{srow['price']}-1", f"=C{srow['pt']}/C{srow['price']}-1", PCT, True)
for k in ("pt", "updown"):
    for col in "BC":
        sm[f"{col}{srow[k]}"].fill = YELLOW

sec("5. Revenue, margins and FCF by year ($mm)")
yrs_show = [2027, 2028, 2029, 2030, 2035, 2040]
put(sm, f"A{r}", "Line", bold=True)
for j, y in enumerate(yrs_show):
    put(sm, f"{L(2+j)}{r}", y, fmt=YRFMT, bold=True)
sm.column_dimensions["G"].width = 12
r += 1
for co in ("CEG", "VST"):
    put(sm, f"A{r}", co, bold=True)
    r += 1
    for k, lab, fmt in (("c_rev", "Revenue", USD), ("c_gm", "Gross margin", USD), ("c_gmpct", "Gross margin %", PCT), ("c_ebitda", "Asset EBITDA", USD),
                        ("c_ebitdapct", "EBITDA margin %", PCT), ("c_assetcash", "Asset cash margin", USD), ("c_fcfadj", "Equity FCF", USD)):
        put(sm, f"A{r}", "  " + lab)
        for j, y in enumerate(yrs_show):
            col = YC[YEARS.index(y)]
            put(sm, f"{L(2+j)}{r}", f"='{co}_Model'!{col}{built[co][3][k]}", GREEN, fmt)
        r += 1
    put(sm, f"A{r}", "  FCF yield on market cap")
    col_m = "B" if co == "CEG" else "C"
    for j, y in enumerate(yrs_show):
        put(sm, f"{L(2+j)}{r}", f"={L(2+j)}{r-1}/${col_m}${srow['mcap']}", fmt=PCT)
    r += 1

snap_row = r + 1
wb.move_sheet("Summary", offset=0)
wb.move_sheet("README", offset=-1)
wb.save(OUT)
print(OUT, "snapshot_row", snap_row)
import json
json.dump({"srow": srow, "snap_row": snap_row}, open(os.path.join(os.path.dirname(OUT), ".xlsx_layout.json"), "w"))

import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L

OUT = sys.argv[1]
flex = dict(gm=0.0, dcai=0.0, fdy=1.0, capex=1.0, cn=0.0)
if len(sys.argv) > 2:
    for kv in sys.argv[2].split(","):
        k, v = kv.split("="); flex[k] = float(v)

wb = Workbook()
F = "Arial"
BLUE = Font(name=F, color="0000FF"); BLACK = Font(name=F); GREEN = Font(name=F, color="008000")
BOLD = Font(name=F, bold=True); TITLE = Font(name=F, bold=True, size=14)
HDR = Font(name=F, bold=True, color="FFFFFF"); HFILL = PatternFill("solid", fgColor="1F3864")
YEL = PatternFill("solid", fgColor="FFFF00"); GREY = PatternFill("solid", fgColor="F2F2F2")
thin = Border(top=Side(style="thin"))
USD = '$#,##0.0;($#,##0.0);-'; USD2 = '$#,##0.00;($#,##0.00);-'; PCT = '0.0%;(0.0%);-'; MULT = '0.0x'; NUM = '#,##0.000'

def put(ws, ref, val, font=BLACK, fmt=None, fill=None, note=None, bold=False):
    c = ws[ref]; c.value = val
    c.font = Font(name=F, bold=bold or font.bold, color=font.color)
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    if note: c.comment = Comment(note, "Model")
    return c

def header(ws, row, labels, start=1):
    for i, t in enumerate(labels):
        c = ws.cell(row=row, column=start + i, value=t); c.font = HDR; c.fill = HFILL
        c.alignment = Alignment(horizontal="center")

YEARS = [str(y) for y in range(2027, 2034)]  # 7 explicit years
YC = [L(3 + i) for i in range(7)]              # C..I

# ---------------- README ----------------
rd = wb.active; rd.title = "README"
lines = [
 ("Intel (INTC) Valuation Model: scenario DCF built from the Sep 2026 product analysis", TITLE),
 ("Companion to research/intel-competitive-analysis-2026-09.md. Values in US$ billions unless stated; per-share values in US$.", BLACK),
 ("", BLACK),
 ("HOW TO USE", BOLD),
 ("1. Edit blue cells only (inputs). Yellow-filled cells are key judgment calls or placeholders you should verify.", BLACK),
 ("2. Black cells are formulas; green cells link to another sheet. Every sheet recalculates when inputs change.", BLACK),
 ("3. Scenario drivers live on 'Inputs' (one row per scenario per driver). Bear/Base/Bull sheets are identical in structure.", BLACK),
 ("4. 'Summary' shows each scenario, the probability-weighted value, and a reverse DCF (what today's price requires).", BLACK),
 ("5. 'Sensitivity' has live WACC x terminal-growth grids, plus a tornado table computed by flexing single inputs.", BLACK),
 ("", BLACK),
 ("METHOD", BOLD),
 ("- Base year 2026E = Q1 + Q2 actuals + Q3 guidance midpoint + Q4 assumption, split by segment using the Q2 2026 mix.", BLACK),
 ("- 2027-2033 explicit forecast by segment: DCAI (server), CCPG (client), external foundry + packaging, other.", BLACK),
 ("- China layer: the China DOMESTIC share of DCAI+CCPG revenue (China-billed share x domestic fraction) erodes at a scenario-specific annual rate", BLACK),
 ("  (xinchuang substitution, Hygon/Loongson/Kunpeng, retaliatory tariffs on US-fabbed chips); the loss is subtracted from revenue at full gross margin.", BLACK),
 ("- Not modeled as a line item: a Taiwan disruption. It would hit Intel too (TSMC makes Lunar/Arrow Lake tiles, part of Panther/Nova Lake).", BLACK),
 ("  Geopolitical second-source demand for US fabs is already inside the external-foundry drivers (Base/Bull), so it isn't added twice.", BLACK),
 ("- Margins on a non-GAAP basis, then stock-based comp deducted as a real cost. Unlevered FCF = NOPAT + D&A - capex - change in NWC.", BLACK),
 ("- Discounted from 30 Sep 2026; Gordon-growth terminal value on normalized 2034 FCF (terminal capex = D&A x ratio).", BLACK),
 ("", BLACK),
 ("HOW SCENARIOS MAP TO THE PRODUCT FINDINGS", BOLD),
 ("Bear: AMD Venice + Arm keep taking server share and the shortage-driven ASP spike reverses; client share erosion continues; no 14A anchor", BLACK),
 ("      customer, external foundry stays mostly Altera + small packaging deals; 18A cost parity slips; gross margin stays near 2026 levels.", BLACK),
 ("Base: Diamond Rapids lands mid-2027 roughly competitive; server share stabilizes; Panther/Nova Lake hold client roughly flat; EMIB-T packaging", BLACK),
 ("      and an Apple 18A-P ramp build external revenue to ~$15bn by 2033; 18A reaches industry-standard yield in 2027; GM recovers to ~50%.", BLACK),
 ("Bull: Intel wins a 14A anchor customer, packaging becomes a multi-billion business, Diamond Rapids wins sockets, AI-driven CPU demand stays", BLACK),
 ("      strong; GM rises to ~57% (still below TSMC's 67.7% and Intel's own 2010s levels ~60%).", BLACK),
 ("", BLACK),
 ("IMPORTANT CAVEATS", BOLD),
 ("- Scenario drivers are MY judgment calls derived from the product evidence, not forecasts I believe will happen. They are labeled hypotheticals.", BLACK),
 ("- Some inputs are placeholders I could not source in this environment (D&A, SBC, non-controlling interests, Q1-26 external foundry revenue).", BLACK),
 ("  They are filled yellow. Replace them from the Q2 2026 10-Q before relying on the output.", BLACK),
 ("- Share count is derived (Q2 non-GAAP net income / non-GAAP EPS + August 2026 offering shares); verify against the 10-Q cover page.", BLACK),
 ("- Scenario assumptions were fixed BEFORE computing values and were not tuned toward the market price.", BLACK),
 ("- Not investment advice.", BLACK),
]
for i, (t, f) in enumerate(lines, 1):
    put(rd, f"A{i}", t, f)
rd.column_dimensions["A"].width = 150

# ---------------- INPUTS ----------------
inp = wb.create_sheet("Inputs")
inp.column_dimensions["A"].width = 52; inp.column_dimensions["B"].width = 14
for col in "CDEFGHI": inp.column_dimensions[col].width = 11
inp.column_dimensions["J"].width = 90
put(inp, "A1", "Inputs (US$ bn unless stated). Blue = input, black = formula, yellow = key judgment / placeholder to verify", TITLE)
N = {}  # name -> absolute ref
def inrow(r, label, val, fmt, name, note, key=False, formula=False):
    put(inp, f"A{r}", label)
    put(inp, f"B{r}", val, BLACK if formula else BLUE, fmt, YEL if key else None)
    put(inp, f"J{r}", note)
    N[name] = f"Inputs!$B${r}"

r = 3; put(inp, f"A{r}", "MARKET DATA & BALANCE SHEET", BOLD); header(inp, r, ["Item", "Value"], 1); inp[f"J{r}"].value = "Source / note"; inp[f"J{r}"].font = HDR; inp[f"J{r}"].fill = HFILL
rows = [
 ("Share price ($)", 123.00, USD2, "price", "Close 25 Sep 2026 per TradingKey (tradingkey.com). Update to current price.", False, False),
 ("Q2 2026 non-GAAP net income attributable to Intel", 2.2, USD, "ni_q2", "Intel Q2 2026 earnings release (23 Jul 2026)", False, False),
 ("Q2 2026 non-GAAP EPS ($)", 0.42, USD2, "eps_q2", "Intel Q2 2026 earnings release (23 Jul 2026)", False, False),
 ("Implied Q2 diluted shares (bn)", "=B5/B6", NUM, "sh_q2", "Derived = NI / EPS (rounding makes this approximate). VERIFY vs 10-Q.", True, True),
 ("Shares issued in Aug 2026 offering incl. greenshoe (bn)", 0.2421, NUM, "sh_off", "210,526,315 + 31,578,947 shares; Intel newsroom, 10-11 Aug 2026", False, False),
 ("Diluted shares for valuation (bn)", "=B7+B8", NUM, "shares", "Excludes any further escrowed government shares not in diluted count; verify.", False, True),
 ("Cash & equivalents (27 Jun 2026)", 12.874, USD, "cash", "Q2 2026 10-Q, via search extract", False, False),
 ("Short-term investments (27 Jun 2026)", 16.853, USD, "sti", "Q2 2026 10-Q, via search extract", False, False),
 ("Net proceeds of Aug 2026 equity offering", 22.62, USD, "raise", "Intel 8-K / newsroom, Aug 2026", False, False),
 ("Total debt (27 Jun 2026)", 48.549, USD, "debt", "Q2 2026 10-Q, via search extract", False, False),
 ("Altera 49% stake (equity-method)", "=8.75*0.49", USD, "altera", "Placeholder: 49% x $8.75bn valuation from the 2025 Silver Lake deal. Verify carrying/fair value.", True, True),
 ("Non-controlling interests (Mobileye minority, SCIP fab partners)", 0.0, USD, "nci", "PLACEHOLDER = 0. Intel's Arizona/Ireland SCIP partners and Mobileye minority holders have material claims. Enter from 10-Q; this overstates value if left at 0.", True, False),
 ("Net cash (debt) incl. offering", "=B10+B11+B12-B13", USD, "netcash", "Formula", False, True),
 ("Market capitalization", "=B4*B9", USD, "mcap", "Formula", False, True),
 ("Enterprise value at market", "=B17-B16-B14+B15", USD, "ev_mkt", "Market cap - net cash - Altera stake + NCI", False, True),
]
r = 4
for lab, v, fmt, nm, note, key, fml in rows:
    inrow(r, lab, v, fmt, nm, note, key, fml); r += 1

r += 1; put(inp, f"A{r}", "2026E BASE YEAR BUILD", BOLD); r += 1
header(inp, r, ["Item", "Value"], 1); r += 1
rows2 = [
 ("Q1 2026 total revenue", 13.6, "q1", "Intel Q1 2026 results (23 Apr 2026)"),
 ("Q2 2026 total revenue", 16.1, "q2", "Intel Q2 2026 results"),
 ("Q3 2026 revenue guidance midpoint", 16.3, "q3", "Guidance $15.8-16.8bn (Q2 call)"),
 ("Q4 2026 revenue assumption", "=B23", "q4", "ASSUMPTION: flat vs Q3 guide midpoint"),
 ("Q1 2026 DCAI revenue", 5.1, "dcai_q1", "Intel Q1 2026 results"),
 ("Q2 2026 DCAI revenue", 6.3, "dcai_q2", "Intel Q2 2026 results"),
 ("Q1 2026 CCPG (client) revenue", 7.7, "ccg_q1", "Intel Q1 2026 results"),
 ("Q2 2026 CCPG (client) revenue", 8.9, "ccg_q2", "Intel Q2 2026 results ($8.88bn)"),
 ("Q1 2026 external foundry revenue", 0.25, "fdy_q1", "PLACEHOLDER (not found); Q2 was $0.293bn"),
 ("Q2 2026 external foundry revenue", 0.293, "fdy_q2", "Intel Q2 2026 10-Q / DCD"),
 ("2026E non-GAAP gross margin", 0.41, "gm26", "Q2 41.8%, Q3 guide 42%; Q1 lower. Assumption."),
 ("2026E non-GAAP operating expenses", 16.5, "opex26", "PLACEHOLDER near FY2025 target of $16.8bn; verify 2026 target"),
 ("2026E stock-based compensation", 2.5, "sbc26", "PLACEHOLDER; verify from 10-Q cash-flow statement"),
 ("2026E depreciation & amortization", 11.0, "da26", "PLACEHOLDER; verify from 10-Q cash-flow statement"),
 ("2026E gross capex", 20.5, "capex26", "Guided 'more than $20bn' (Q2 call)"),
]
start2 = r
for lab, v, nm, note in rows2:
    key = "PLACEHOLDER" in note or "ASSUMPTION" in note
    fmt = PCT if nm == "gm26" else USD
    inrow(r, lab, v, fmt, nm, note, key, isinstance(v, str)); r += 1
# fix q4 formula to reference q3 row
inp[N["q4"].split("!")[1].replace("$", "")].value = "=" + N["q3"].split("!")[1].replace("$", "")
# derived 2026E segments
def ref(n): return N[n].split("!")[1]
derived = [
 ("2026E total revenue", f"={ref('q1')}+{ref('q2')}+{ref('q3')}+{ref('q4')}", "rev26", "Formula"),
 ("2026E DCAI revenue", f"={ref('dcai_q1')}+{ref('dcai_q2')}+({ref('q3')}+{ref('q4')})*{ref('dcai_q2')}/{ref('q2')}", "dcai26", "H2 allocated at Q2 mix"),
 ("2026E CCPG revenue", f"={ref('ccg_q1')}+{ref('ccg_q2')}+({ref('q3')}+{ref('q4')})*{ref('ccg_q2')}/{ref('q2')}", "ccg26", "H2 allocated at Q2 mix"),
 ("2026E external foundry revenue", f"={ref('fdy_q1')}+{ref('fdy_q2')}+({ref('q3')}+{ref('q4')})*{ref('fdy_q2')}/{ref('q2')}", "fdy26", "H2 allocated at Q2 mix"),
 ("2026E other revenue (Mobileye etc.)", "=B{a}-B{b}-B{c}-B{d}", "oth26", "Residual"),
]
for lab, f, nm, note in derived:
    inrow(r, lab, f, USD, nm, note, False, True); r += 1
inp[ref("oth26")].value = f"={ref('rev26')}-{ref('dcai26')}-{ref('ccg26')}-{ref('fdy26')}"

r += 1; put(inp, f"A{r}", "CHINA EXPOSURE", BOLD); r += 1
header(inp, r, ["Item", "Value"], 1); r += 1
for lab, v, fmt, nm, note, key in [
 ("FY2024 China (incl. HK) billed revenue", 15.53, USD, "cn24", "Intel FY2024 10-K geographic note (billing location); 29% of revenue", False),
 ("FY2025 total revenue", 52.9, USD, "rev25", "Intel Q4/FY2025 results (22 Jan 2026)", False),
 ("FY2025 US + Taiwan + Singapore + Other billed revenue", "=15.76+7.67+9.54+7.20", USD, "noncn25", "FY2025 10-K geographic note via search extract (US 15.76, Taiwan 7.67, Singapore 9.54, Other 7.20); VERIFY", True),
 ("FY2025 China billed revenue (derived residual)", "=B{r}", USD, "cn25", "Derived = total - other regions. Not read directly from the 10-K; VERIFY", True),
 ("China billed share of revenue (FY2025)", "=0", PCT, "cnshare", "Formula", False),
 ("Share of China-billed revenue that is China DOMESTIC end-demand", 0.50, PCT, "cndom", "JUDGMENT. Billing location includes PCs/servers assembled in China for export (Lenovo, ODMs), which domestic substitution does not hit. No source splits this; test 0.35-0.70", True),
]:
    inrow(r, lab, v, fmt, nm, note, key, isinstance(v, str)); r += 1
inp[ref("cn25")].value = f"={ref('rev25')}-{ref('noncn25')}"
inp[ref("cnshare")].value = f"=IF({ref('rev25')}=0,0,{ref('cn25')}/{ref('rev25')})"
put(inp, f"J{r}", "China levers: xinchuang bans Intel/AMD in government PCs & servers (Mar 2024); Hygon+Zhaoxin ~15-20% of China server CPUs, Kunpeng 8-12% (2026 est.); Loongson 1M desktop CPUs; China customs treats WAFER-FAB location as origin -> US-fabbed chips (Intel) face retaliatory tariffs, TSMC-made rivals don't; US-China truce expires 10 Nov 2026.")
r += 1

r += 1; put(inp, f"A{r}", "VALUATION PARAMETERS", BOLD); r += 1
header(inp, r, ["Item", "Value"], 1); r += 1
rows3 = [
 ("Risk-free rate", 0.043, PCT, "rf", "ASSUMPTION: ~10y UST; update", True),
 ("Equity risk premium", 0.05, PCT, "erp", "ASSUMPTION", True),
 ("Equity beta", 1.3, '0.00', "beta", "ASSUMPTION: semis/turnaround risk", True),
 ("Cost of equity", f"={'B'}{{rf}}+{'B'}{{beta}}*{'B'}{{erp}}", PCT, "ke", "CAPM", False),
 ("Pre-tax cost of debt", 0.055, PCT, "kd", "ASSUMPTION", False),
 ("Tax rate", 0.15, PCT, "tax", "ASSUMPTION: normalized cash tax; Intel has large NOLs near term", True),
 ("Debt weight", "=debt/(debt+mcap)", PCT, "wd", "Market-value weights", False),
 ("WACC", "=wacc", PCT, "wacc", "Formula", False),
 ("Terminal FCF growth", 0.03, PCT, "g", "ASSUMPTION", True),
 ("Terminal capex / D&A ratio", 1.15, MULT, "tcapex", "ASSUMPTION: maintenance + modest growth capex in steady state", False),
 ("Net working capital, % of revenue change", 0.10, PCT, "nwc", "ASSUMPTION", False),
 ("D&A catch-up toward prior-year capex (per year)", 0.15, PCT, "dacu", "D&A(t) = D&A(t-1) + x * (capex(t-1) - D&A(t-1))", False),
 ("Discount period offset (valuation date 30 Sep 2026)", 0.25, '0.00', "off", "End-2027 cash flow is 1.25 years out", False),
 ("Other revenue growth (all scenarios)", 0.05, PCT, "othg", "ASSUMPTION", False),
]
for lab, v, fmt, nm, note, key in rows3:
    inrow(r, lab, 0, fmt, nm, note, key, isinstance(v, str))
    if not isinstance(v, str): inp[f"B{r}"].value = v
    r += 1
inp[ref("ke")].value = f"={ref('rf')}+{ref('beta')}*{ref('erp')}"
inp[ref("wd")].value = f"={ref('debt')}/({ref('debt')}+{ref('mcap')})"
inp[ref("wacc")].value = f"=(1-{ref('wd')})*{ref('ke')}+{ref('wd')}*{ref('kd')}*(1-{ref('tax')})"

r += 1; put(inp, f"A{r}", "FLEX LEVERS (for sensitivity; leave at defaults 0 / 0 / 1 / 1)", BOLD); r += 1
for lab, v, fmt, nm, note in [
 ("Gross margin shift, all years, all scenarios (pts)", flex["gm"], PCT, "fx_gm", "e.g. -0.03 = 3 pts lower"),
 ("DCAI growth shift per year (pts)", flex["dcai"], PCT, "fx_dcai", "e.g. 0.03 = +3 pts/yr"),
 ("External foundry revenue multiplier", flex["fdy"], MULT, "fx_fdy", "e.g. 0.5 = half"),
 ("Capex multiplier (2027+)", flex["capex"], MULT, "fx_capex", "e.g. 1.15 = 15% more"),
 ("China domestic revenue change shift per year (pts)", flex["cn"], PCT, "fx_cn", "e.g. -0.10 = 10 pts faster erosion every year"),
]:
    inrow(r, lab, v, fmt, nm, note); r += 1

r += 1; put(inp, f"A{r}", "SCENARIO WEIGHTS (judgment; must sum to 100%)", BOLD); r += 1
header(inp, r, ["Scenario", "Weight"], 1); r += 1
for lab, v, nm, note in [("Bear", 0.30, "w_bear", "My judgment; edit"), ("Base", 0.45, "w_base", "My judgment; edit"), ("Bull", 0.25, "w_bull", "My judgment; edit")]:
    inrow(r, lab, v, PCT, nm, note, True); r += 1
inrow(r, "Check: sum of weights", f"={ref('w_bear')}+{ref('w_base')}+{ref('w_bull')}", PCT, "w_sum", "Must equal 100%", False, True); r += 1

# driver tables
r += 1; put(inp, f"A{r}", "SCENARIO DRIVERS (2027-2033). Hypothetical paths tied to product evidence; NOT forecasts I endorse", BOLD); r += 1
drivers = {
 "DCAI revenue growth": ("dcai", PCT, {
   "Bear": [-0.12, -0.03, 0.00, 0.02, 0.02, 0.02, 0.02],
   "Base": [0.04, 0.05, 0.06, 0.06, 0.05, 0.04, 0.04],
   "Bull": [0.15, 0.12, 0.10, 0.08, 0.07, 0.06, 0.05]},
   "Bear: shortage ASP spike (+48% YoY in Q2-26) reverses as Venice ships Q4-26 and Diamond Rapids slips; Bull: DMR competitive + agentic CPU demand persists"),
 "CCPG (client) revenue growth": ("ccg", PCT, {
   "Bear": [-0.08, -0.04, -0.03, -0.02, -0.02, -0.01, -0.01],
   "Base": [-0.02, 0.01, 0.02, 0.02, 0.02, 0.02, 0.02],
   "Bull": [0.04, 0.04, 0.03, 0.03, 0.03, 0.03, 0.03]},
   "Mercury: AMD client share 30.3% Q2-26 and rising; Arm WoA ~3%. Bull assumes Nova Lake + NVIDIA x86 RTX SoCs hold share"),
 "External foundry + packaging revenue ($bn)": ("fdy", USD, {
   "Bear": [1.3, 1.6, 2.0, 2.5, 3.0, 3.5, 4.0],
   "Base": [1.8, 3.5, 6.0, 8.5, 11.0, 13.0, 15.0],
   "Bull": [3.0, 7.0, 12.0, 18.0, 24.0, 29.0, 33.0]},
   "Today ~$1.2bn annualized, mostly Altera. Base = EMIB-T packaging + Apple 18A-P ramp. Bull = 14A anchor customer. All ABOVE Intel's historical base rate"),
 "Non-GAAP gross margin": ("gm", PCT, {
   "Bear": [0.39, 0.39, 0.40, 0.40, 0.40, 0.40, 0.40],
   "Base": [0.43, 0.45, 0.47, 0.48, 0.49, 0.50, 0.50],
   "Bull": [0.46, 0.49, 0.52, 0.54, 0.55, 0.56, 0.57]},
   "Q2-26 41.8%. Compare TSMC 67.7% (Q2-26), Intel ~60% in the 2010s. Base requires 18A industry-standard yields in 2027 (CFO's own timeline)"),
 "Non-GAAP opex growth": ("opx", PCT, {
   "Bear": [0.01]*7, "Base": [0.03]*7, "Bull": [0.04]*7},
   "Bull spends more (AI GPU, foundry customer support)"),
 "China domestic end-demand revenue change per year": ("cn", PCT, {
   "Bear": [-0.20, -0.15, -0.15, -0.10, -0.10, -0.10, -0.10],
   "Base": [-0.08, -0.08, -0.07, -0.07, -0.06, -0.06, -0.05],
   "Bull": [-0.03, -0.03, -0.03, -0.03, -0.03, -0.03, -0.03]},
   "Bear: truce lapses Nov-26, xinchuang spreads from government to SOEs/finance/telecom, and origin rule hits 18A (US-fabbed) parts. Base: steady localization. Bull: truce holds, erosion slow"),
 "Gross capex ($bn)": ("cpx", USD, {
   "Bear": [20, 18, 16, 15, 15, 15, 15],
   "Base": [26, 26, 24, 22, 20, 20, 20],
   "Bull": [30, 32, 30, 28, 26, 26, 26]},
   "2027 guided 'significantly above' 2026 (>$20bn). Gross; excludes SCIP partner/government offsets (conservative)"),
}
DR = {}  # (key, scen) -> row
for title, (key, fmt, vals, note) in drivers.items():
    put(inp, f"A{r}", title, BOLD); header(inp, r, [""] + YEARS, 2); inp[f"B{r}"].value = "Scenario"; r += 1
    for sc in ["Bear", "Base", "Bull"]:
        put(inp, f"A{r}", f"  {title} - {sc}"); put(inp, f"B{r}", sc)
        for i, v in enumerate(vals[sc]):
            put(inp, f"{YC[i]}{r}", v, BLUE, fmt, YEL)
        DR[(key, sc)] = r; r += 1
    put(inp, f"A{r}", "  Rationale: " + note, Font(name=F, italic=True)); r += 2

# ---------------- SCENARIO SHEETS ----------------
def I(n): return N[n]
def scen(sc):
    ws = wb.create_sheet(sc)
    ws.column_dimensions["A"].width = 46
    for col in "BCDEFGHIJ": ws.column_dimensions[col].width = 12
    put(ws, "A1", f"{sc} scenario: US$ bn", TITLE)
    header(ws, 3, ["Line item", "2026E"] + YEARS + ["Terminal"], 1)
    lab = {}
    items = ["period", "dcai", "dcai_g", "ccg", "ccg_g", "fdy", "oth", "chx", "chf", "chl", "rev", "rev_g", "gm", "gp", "opex", "oi", "oi_m",
             "sbc", "ebit", "tax", "nopat", "da", "capex", "nwc", "fcf", "fcf_m", "df", "pv"]
    names = {"period": "Discount period (yrs from 30 Sep 2026)", "dcai": "DCAI (server/data center) revenue", "dcai_g": "  growth",
             "ccg": "CCPG (client) revenue", "ccg_g": "  growth", "fdy": "External foundry + packaging revenue", "oth": "Other revenue (Mobileye etc.)",
             "chx": "China domestic-demand revenue exposed (pre-erosion)", "chf": "  China erosion index (2026 = 1.00)",
             "chl": "Less: China localization / trade revenue loss", "rev": "Total revenue", "rev_g": "  growth", "gm": "Non-GAAP gross margin", "gp": "Gross profit", "opex": "Non-GAAP opex",
             "oi": "Non-GAAP operating income", "oi_m": "  operating margin", "sbc": "Stock-based compensation", "ebit": "EBIT after SBC",
             "tax": "Taxes on EBIT", "nopat": "NOPAT", "da": "D&A", "capex": "Capex", "nwc": "Change in net working capital",
             "fcf": "Unlevered free cash flow", "fcf_m": "  FCF margin", "df": "Discount factor", "pv": "PV of FCF"}
    for i, k in enumerate(items):
        rr = 4 + i; lab[k] = rr
        put(ws, f"A{rr}", names[k], BOLD if k in ("rev", "oi", "fcf") else BLACK)
    R = lab
    # 2026E column B
    put(ws, f"B{R['dcai']}", f"={I('dcai26')}", GREEN, USD)
    put(ws, f"B{R['ccg']}", f"={I('ccg26')}", GREEN, USD)
    put(ws, f"B{R['fdy']}", f"={I('fdy26')}", GREEN, USD)
    put(ws, f"B{R['oth']}", f"={I('oth26')}", GREEN, USD)
    put(ws, f"B{R['chx']}", f"=(B{R['dcai']}+B{R['ccg']})*{I('cnshare')}*{I('cndom')}", BLACK, USD)
    put(ws, f"B{R['chf']}", 1, BLUE, '0.00')
    put(ws, f"B{R['chl']}", f"=-B{R['chx']}*(1-B{R['chf']})", BLACK, USD)
    put(ws, f"B{R['rev']}", f"=SUM(B{R['dcai']},B{R['ccg']},B{R['fdy']},B{R['oth']},B{R['chl']})", BLACK, USD, bold=True)
    put(ws, f"B{R['gm']}", f"={I('gm26')}+{I('fx_gm')}", GREEN, PCT)
    put(ws, f"B{R['gp']}", f"=B{R['rev']}*B{R['gm']}", BLACK, USD)
    put(ws, f"B{R['opex']}", f"={I('opex26')}", GREEN, USD)
    put(ws, f"B{R['oi']}", f"=B{R['gp']}-B{R['opex']}", BLACK, USD, bold=True)
    put(ws, f"B{R['oi_m']}", f"=IF(B{R['rev']}=0,0,B{R['oi']}/B{R['rev']})", BLACK, PCT)
    put(ws, f"B{R['sbc']}", f"={I('sbc26')}", GREEN, USD)
    put(ws, f"B{R['da']}", f"={I('da26')}", GREEN, USD)
    put(ws, f"B{R['capex']}", f"={I('capex26')}", GREEN, USD)
    for i, c in enumerate(YC):
        p = "B" if i == 0 else YC[i - 1]
        dcol = YC[i]  # driver column in Inputs
        put(ws, f"{c}{R['period']}", f"={I('off')}+{i+1}", BLACK, '0.00')
        put(ws, f"{c}{R['dcai_g']}", f"=Inputs!{dcol}{DR[('dcai', sc)]}+{I('fx_dcai')}", GREEN, PCT)
        put(ws, f"{c}{R['dcai']}", f"={p}{R['dcai']}*(1+{c}{R['dcai_g']})", BLACK, USD)
        put(ws, f"{c}{R['ccg_g']}", f"=Inputs!{dcol}{DR[('ccg', sc)]}", GREEN, PCT)
        put(ws, f"{c}{R['ccg']}", f"={p}{R['ccg']}*(1+{c}{R['ccg_g']})", BLACK, USD)
        put(ws, f"{c}{R['fdy']}", f"=Inputs!{dcol}{DR[('fdy', sc)]}*{I('fx_fdy')}", GREEN, USD)
        put(ws, f"{c}{R['oth']}", f"={p}{R['oth']}*(1+{I('othg')})", BLACK, USD)
        put(ws, f"{c}{R['chx']}", f"=({c}{R['dcai']}+{c}{R['ccg']})*{I('cnshare')}*{I('cndom')}", BLACK, USD)
        put(ws, f"{c}{R['chf']}", f"=MAX(0,{p}{R['chf']}*(1+Inputs!{dcol}{DR[('cn', sc)]}+{I('fx_cn')}))", BLACK, '0.00')
        put(ws, f"{c}{R['chl']}", f"=-{c}{R['chx']}*(1-{c}{R['chf']})", BLACK, USD)
        put(ws, f"{c}{R['rev']}", f"=SUM({c}{R['dcai']},{c}{R['ccg']},{c}{R['fdy']},{c}{R['oth']},{c}{R['chl']})", BLACK, USD, bold=True)
        put(ws, f"{c}{R['rev_g']}", f"=IF({p}{R['rev']}=0,0,{c}{R['rev']}/{p}{R['rev']}-1)", BLACK, PCT)
        put(ws, f"{c}{R['gm']}", f"=Inputs!{dcol}{DR[('gm', sc)]}+{I('fx_gm')}", GREEN, PCT)
        put(ws, f"{c}{R['gp']}", f"={c}{R['rev']}*{c}{R['gm']}", BLACK, USD)
        put(ws, f"{c}{R['opex']}", f"={p}{R['opex']}*(1+Inputs!{dcol}{DR[('opx', sc)]})", BLACK, USD)
        put(ws, f"{c}{R['oi']}", f"={c}{R['gp']}-{c}{R['opex']}", BLACK, USD, bold=True)
        put(ws, f"{c}{R['oi_m']}", f"=IF({c}{R['rev']}=0,0,{c}{R['oi']}/{c}{R['rev']})", BLACK, PCT)
        put(ws, f"{c}{R['sbc']}", f"={p}{R['sbc']}*(1+Inputs!{dcol}{DR[('opx', sc)]})", BLACK, USD)
        put(ws, f"{c}{R['ebit']}", f"={c}{R['oi']}-{c}{R['sbc']}", BLACK, USD)
        put(ws, f"{c}{R['tax']}", f"=MAX(0,{c}{R['ebit']})*{I('tax')}", BLACK, USD)
        put(ws, f"{c}{R['nopat']}", f"={c}{R['ebit']}-{c}{R['tax']}", BLACK, USD)
        put(ws, f"{c}{R['capex']}", f"=Inputs!{dcol}{DR[('cpx', sc)]}*{I('fx_capex')}", GREEN, USD)
        put(ws, f"{c}{R['da']}", f"={p}{R['da']}+{I('dacu')}*({p}{R['capex']}-{p}{R['da']})", BLACK, USD)
        put(ws, f"{c}{R['nwc']}", f"={I('nwc')}*({c}{R['rev']}-{p}{R['rev']})", BLACK, USD)
        put(ws, f"{c}{R['fcf']}", f"={c}{R['nopat']}+{c}{R['da']}-{c}{R['capex']}-{c}{R['nwc']}", BLACK, USD, bold=True)
        put(ws, f"{c}{R['fcf_m']}", f"=IF({c}{R['rev']}=0,0,{c}{R['fcf']}/{c}{R['rev']})", BLACK, PCT)
        put(ws, f"{c}{R['df']}", f"=1/(1+{I('wacc')})^{c}{R['period']}", BLACK, '0.000')
        put(ws, f"{c}{R['pv']}", f"={c}{R['fcf']}*{c}{R['df']}", BLACK, USD)
    # terminal column J (normalized 2034)
    g = I('g'); last = "I"
    put(ws, f"J{R['rev']}", f"={last}{R['rev']}*(1+{g})", BLACK, USD, bold=True)
    put(ws, f"J{R['nopat']}", f"={last}{R['nopat']}*(1+{g})", BLACK, USD)
    put(ws, f"J{R['da']}", f"={last}{R['da']}*(1+{g})", BLACK, USD)
    put(ws, f"J{R['capex']}", f"=J{R['da']}*{I('tcapex')}", BLACK, USD)
    put(ws, f"J{R['nwc']}", f"={I('nwc')}*(J{R['rev']}-{last}{R['rev']})", BLACK, USD)
    put(ws, f"J{R['fcf']}", f"=J{R['nopat']}+J{R['da']}-J{R['capex']}-J{R['nwc']}", BLACK, USD, bold=True)
    put(ws, f"J{R['fcf_m']}", f"=IF(J{R['rev']}=0,0,J{R['fcf']}/J{R['rev']})", BLACK, PCT)
    for k in ("rev", "oi", "fcf"):
        for col in "ABCDEFGHIJ": ws[f"{col}{R[k]}"].border = thin
    # valuation block
    v0 = R["pv"] + 3
    put(ws, f"A{v0-1}", "VALUATION (value per share floored at $0: equity has limited liability)", BOLD)
    V = {}
    vrows = [
     ("pvsum", "Sum of PV of explicit FCF (2027-2033)", f"=SUM(C{R['pv']}:I{R['pv']})", USD),
     ("tv", "Terminal value at end-2033", f"=IF({I('wacc')}<={g},0,J{R['fcf']}/({I('wacc')}-{g}))", USD),
     ("pvtv", "PV of terminal value", f"=TV*I{R['df']}", USD),
     ("ev", "Enterprise value", "=PVSUM+PVTV", USD),
     ("netcash", "+ Net cash (incl. Aug-26 raise)", f"={I('netcash')}", USD),
     ("altera", "+ Altera stake", f"={I('altera')}", USD),
     ("nci", "- Non-controlling interests", f"=-{I('nci')}", USD),
     ("eq", "Equity value", "=EV+NETCASH+ALTERA+NCI", USD),
     ("sh", "Diluted shares (bn)", f"={I('shares')}", NUM),
     ("vps", "Value per share ($)", "=IF(SH=0,0,MAX(0,EQ/SH))", USD2),
     ("px", "Current share price ($)", f"={I('price')}", USD2),
     ("up", "Upside / (downside) vs price", "=IF(PX=0,0,VPS/PX-1)", PCT),
     ("tvpct", "Terminal value as % of EV", "=IF(EV=0,0,PVTV/EV)", PCT),
     ("pe27", "Implied P/E on 2027 NOPAT/share at value", f"=IF(C{R['nopat']}<=0,0,VPS/(C{R['nopat']}/SH))", MULT),
     ("mktpe27", "Market P/E on 2027 NOPAT/share", f"=IF(C{R['nopat']}<=0,0,PX/(C{R['nopat']}/SH))", MULT),
    ]
    for i, (k, _, _, _) in enumerate(vrows): V[k] = v0 + i
    for i, (k, lbl, f, fmt) in enumerate(vrows):
        rr = v0 + i
        import re as _re
        f = _re.sub(r"\b(PVSUM|PVTV|TV|EV|NETCASH|ALTERA|NCI|EQ|SH|VPS|PX|UP|TVPCT)\b", lambda m: f"B{V[m.group(1).lower()]}", f)
        font = GREEN if f.startswith("=Inputs") or f.startswith("=-Inputs") else BLACK
        put(ws, f"A{rr}", lbl, BOLD if k in ("ev", "eq", "vps") else BLACK)
        c = put(ws, f"B{rr}", f, font, fmt, bold=k in ("vps",))
        if k == "vps": c.fill = GREY
    ws.freeze_panes = "B4"
    return R, V

SR = {}
for sc in ["Bear", "Base", "Bull"]:
    SR[sc] = scen(sc)
R, V = SR["Base"]

# ---------------- SUMMARY ----------------
sm = wb.create_sheet("Summary", 1)
sm.column_dimensions["A"].width = 58
for col in "BCDE": sm.column_dimensions[col].width = 15
sm.column_dimensions["F"].width = 70
put(sm, "A1", "Summary: scenario values vs market price (US$ bn unless stated)", TITLE)
header(sm, 3, ["Metric", "Bear", "Base", "Bull", "Prob-weighted"], 1)
metrics = [
 ("Scenario weight", None, "w", PCT),
 ("Value per share ($)", "vps", None, USD2),
 ("Upside / (downside) vs current price", "up", None, PCT),
 ("Enterprise value", "ev", None, USD),
 ("Terminal value % of EV", "tvpct", None, PCT),
 ("2030 revenue", ("row", "rev", "F"), None, USD),
 ("2033 revenue", ("row", "rev", "I"), None, USD),
 ("2030 non-GAAP gross margin", ("row", "gm", "F"), None, PCT),
 ("2030 non-GAAP operating margin", ("row", "oi_m", "F"), None, PCT),
 ("2033 external foundry + packaging revenue", ("row", "fdy", "I"), None, USD),
 ("2033 revenue lost to China localization / trade", ("row", "chl", "I"), None, USD),
 ("2030 unlevered FCF", ("row", "fcf", "F"), None, USD),
 ("Cumulative FCF 2027-2033", ("sum", "fcf"), None, USD),
]
wref = {"Bear": I('w_bear'), "Base": I('w_base'), "Bull": I('w_bull')}
for i, (lbl, key, special, fmt) in enumerate(metrics):
    rr = 4 + i
    put(sm, f"A{rr}", lbl, BOLD if key == "vps" else BLACK)
    for j, sc in enumerate(["Bear", "Base", "Bull"]):
        col = "BCD"[j]; Rs, Vs = SR[sc]
        if special == "w": f = f"={wref[sc]}"
        elif isinstance(key, tuple) and key[0] == "row": f = f"={sc}!{key[2]}{Rs[key[1]]}"
        elif isinstance(key, tuple): f = f"=SUM({sc}!C{Rs[key[1]]}:I{Rs[key[1]]})"
        else: f = f"={sc}!B{Vs[key]}"
        put(sm, f"{col}{rr}", f, GREEN, fmt)
    if key == "vps":
        put(sm, f"E{rr}", "=SUMPRODUCT(B4:D4,B5:D5)", BLACK, USD2, fill=GREY, bold=True)
        VPSW = f"E{rr}"
    elif key == "up":
        put(sm, f"E{rr}", f"=IF({I('price')}=0,0,{VPSW}/{I('price')}-1)", BLACK, PCT, bold=True)
put(sm, "F4", "Weights are my judgment (Inputs sheet); they must sum to 100%")
put(sm, "F5", "Check weights sum:"); put(sm, "F6", f'=IF(ABS({I("w_sum")}-1)<0.0001,"OK","WEIGHTS DO NOT SUM TO 100%")')

rr = 4 + len(metrics) + 1
put(sm, f"A{rr}", "MARKET SNAPSHOT", BOLD); rr += 1
for lbl, f, fmt in [("Share price ($)", f"={I('price')}", USD2), ("Diluted shares (bn)", f"={I('shares')}", NUM),
                    ("Market cap", f"={I('mcap')}", USD), ("Enterprise value at market", f"={I('ev_mkt')}", USD),
                    ("WACC", f"={I('wacc')}", PCT), ("Terminal growth", f"={I('g')}", PCT)]:
    put(sm, f"A{rr}", lbl); put(sm, f"B{rr}", f, GREEN, fmt); rr += 1

rr += 1; put(sm, f"A{rr}", "REVERSE DCF: WHAT DOES TODAY'S PRICE REQUIRE? (holding Base explicit-period cash flows)", BOLD); rr += 1
rv0 = rr
rev = [
 ("EV at market", f"={I('ev_mkt')}", USD),
 ("PV of Base explicit FCF 2027-2033", f"=Base!B{V['pvsum']}", USD),
 ("PV of terminal value required", f"=B{rv0}-B{rv0+1}", USD),
 ("Required normalized 2034 FCF", f"=B{rv0+2}/Base!I{R['df']}*({I('wacc')}-{I('g')})", USD),
 ("Base-case normalized 2034 FCF", f"=Base!J{R['fcf']}", USD),
 ("Required / Base terminal FCF", f"=IF(B{rv0+4}=0,0,B{rv0+3}/B{rv0+4})", MULT),
 ("Base 2034 revenue", f"=Base!J{R['rev']}", USD),
 ("Required 2034 FCF margin on Base revenue", f"=IF(B{rv0+6}=0,0,B{rv0+3}/B{rv0+6})", PCT),
 ("Base 2034 FCF margin", f"=Base!J{R['fcf_m']}", PCT),
 ("Reference: TSMC Q2-26 operating margin (not FCF margin)", 0.603, PCT),
 ("Market-implied perpetual FCF from today (EV x (WACC - g))", f"={I('ev_mkt')}*({I('wacc')}-{I('g')})", USD),
 ("Base 2027 FCF, for comparison", f"=Base!C{R['fcf']}", USD),
 ("Required 2034 revenue if FCF margin = 20%", f"=B{rv0+3}/0.20", USD),
 ("Required 2034 revenue if FCF margin = 25%", f"=B{rv0+3}/0.25", USD),
 ("Required 2034 revenue if FCF margin = 30%", f"=B{rv0+3}/0.30", USD),
 ("Bull-case 2034 revenue, for comparison", f"=Bull!J{SR['Bull'][0]['rev']}", USD),
 ("Bull-case 2034 FCF margin, for comparison", f"=Bull!J{SR['Bull'][0]['fcf_m']}", PCT),
]
for i, (lbl, f, fmt) in enumerate(rev):
    put(sm, f"A{rv0+i}", lbl)
    put(sm, f"B{rv0+i}", f, BLUE if not isinstance(f, str) else (GREEN if "Base!" in f and "-" not in f[1:] else BLACK), fmt)
put(sm, f"F{rv0+9}", "Source: TSMC Q2 2026 results (60.3% op margin). Context only")
put(sm, f"F{rv0+3}", "FCF in 2034 that the Base explicit years plus the market price imply")
rr = rv0 + len(rev) + 1
put(sm, f"A{rr}", "HOW TO READ THIS", BOLD); rr += 1
for t in [
 "- 'If the scenario is true, the math says X.' None of the scenarios is a forecast I endorse. The weights express my judgment of plausibility from the product evidence.",
 "- The reverse DCF shows how much more (or less) terminal cash flow than my Base case the current price needs. A ratio well above 1.0x means the price embeds",
 "  assumptions closer to the Bull case: a 14A anchor customer, packaging at scale, sustained CPU pricing and gross margins in the mid-50s.",
 "- Terminal value is a large share of EV in every scenario. Small changes in WACC or terminal growth move the answer a lot: see the Sensitivity sheet.",
]:
    put(sm, f"A{rr}", t); rr += 1

# ---------------- SENSITIVITY ----------------
se = wb.create_sheet("Sensitivity", 2)
se.column_dimensions["A"].width = 30
for col in "BCDEFGH": se.column_dimensions[col].width = 12
put(se, "A1", "Sensitivity: value per share ($)", TITLE)
waccs = [0.085, 0.095, 0.105, 0.115, 0.125]; gs = [0.02, 0.025, 0.03, 0.035, 0.04]
def grid(top, sc, title):
    Rs, Vs = SR[sc]
    put(se, f"A{top}", title, BOLD)
    put(se, f"A{top+1}", "WACC (down) / terminal g (across)")
    for j, gv in enumerate(gs): put(se, f"{L(2+j)}{top+1}", gv, BLUE, PCT)
    for i, wv in enumerate(waccs):
        rr = top + 2 + i
        put(se, f"A{rr}", wv, BLUE, PCT)
        for j in range(len(gs)):
            cc = L(2 + j); w = f"$A{rr}"; gg = f"{cc}${top+1}"
            fcf = f"{sc}!$C${Rs['fcf']}:$I${Rs['fcf']}"; per = f"{sc}!$C${Rs['period']}:$I${Rs['period']}"
            tv = (f"({sc}!$I${Rs['nopat']}*(1+{gg})+{sc}!$I${Rs['da']}*(1+{gg})*(1-{I('tcapex')})"
                  f"-{I('nwc')}*{sc}!$I${Rs['rev']}*{gg})/({w}-{gg})/(1+{w})^{sc}!$I${Rs['period']}")
            f = (f"=IF({w}<={gg},0,(SUMPRODUCT({fcf},1/(1+{w})^{per})+{tv}"
                 f"+{I('netcash')}+{I('altera')}-{I('nci')})/{I('shares')})")
            f = "=MAX(0," + f[1:] + ")"
            put(se, f"{cc}{rr}", f, BLACK, USD2)
    return top + 2 + len(waccs) + 1
t = 3
for sc in ["Bear", "Base", "Bull"]:
    t = grid(t, sc, f"{sc} scenario: value/share vs WACC and terminal growth (live formulas)")
put(se, f"A{t}", "Current price ($)"); put(se, f"B{t}", f"={I('price')}", GREEN, USD2)
TORNADO_ROW = t + 2
put(se, f"A{TORNADO_ROW}", "TORNADO (Base scenario): one input flexed at a time", BOLD)
wb._tornado_row = TORNADO_ROW
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print(TORNADO_ROW)
